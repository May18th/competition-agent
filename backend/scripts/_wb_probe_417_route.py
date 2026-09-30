# -*- coding: utf-8 -*-
"""417轮探针：走真实 /api/export_word 路由，看产物有没有封面+目录"""
import os, sys, json
os.chdir(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.getcwd())

md = []
for i, ch in enumerate(['项目背景与痛点', '问题量化', '解决方案', '技术路线',
                        '市场与商业模式', '团队与分工', '财务测算', '风险控制',
                        '社会价值', '实施计划'], 1):
    cn = '一二三四五六七八九十'[i - 1]
    md.append('## %s、%s' % (cn, ch))
    md.append(('本章围绕%s展开说明，结合现场实测与调研数据给出可执行的安排与测算依据。' % ch) * 10)
text = '\n\n'.join(md)

import app as A
cli = A.app.test_client()
payload = {
    "title": "作品说明书",
    "text": text,
    "doc_type": "default",
    "competition_name": "iCAN大学生创新创业大赛",
    "project_name": "智能危化品试剂柜",
    "school": "某某大学", "team": "启航团队", "advisor": "张老师",
}
r = cli.post('/api/export_word', json=payload)
print('status=', r.status_code)
if r.status_code != 200:
    try:
        print(r.get_json())
    except Exception:
        print(r.data[:500])
    sys.exit(0)
open('_probe_route.docx', 'wb').write(r.data)
print('bytes=', len(r.data))
import docx
d = docx.Document('_probe_route.docx')
print('paras=', len(d.paragraphs))
for p in d.paragraphs[:12]:
    print(' ', repr(p.style.name), repr(p.text[:40]))
print('has_toc_word=', any('目' in p.text and '录' in p.text for p in d.paragraphs))
