#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""417 轮探针：大纲 N 页 -> PPT 实际 N 页（真出 pptx，不 mock 结构链路，只 mock LLM 资产）"""
import os
import sys

import rich_pipeline as rp
from pptx import Presentation
from pptx_builder import build_deck
from export_guard import check_ppt_matches_outline

PAGES = [
    "项目封面", "痛点与问题量化", "现有方案不足", "我们的解决方案",
    "核心技术架构", "功能演示", "试点实测数据", "成本与定价",
    "商业模式", "市场空间测算", "竞品对比", "实施路线图",
    "团队分工", "风险与应对", "财务预测", "愿景与致谢",
]

# 模拟 competition_agents 生成的 ppt_outline markdown 格式
OUTLINE = "## PPT 大纲\n\n"
for i, t in enumerate(PAGES, 1):
    OUTLINE += "## %d. %s\n\n- 要点一：%s相关关键事实\n- 要点二：%s的量化支撑\n\n" % (i, t, t, t)


def main():
    print("大纲页数：", len(PAGES))

    # 1) 大纲解析
    deck = rp._outline_deck(OUTLINE)
    slides = deck.get("slides") or []
    print("_outline_deck 解析：%d 页，首页 type=%s" % (len(slides), slides[0].get("type")))

    # 2) build_rich_assets（mock 掉 LLM 资产设计，结构链路真跑）
    def _mock_assets(result, competition, idea, kb_data=""):
        return {"charts": [], "tables": []}
    rp._gen_assets = _mock_assets
    result = {"ppt_outline": OUTLINE, "one_liner": "校园快递智能柜", "user_data": ""}
    out = rp.build_rich_assets(result, competition="iCAN", idea="校园快递智能柜", theme="tech")
    deck2 = out.get("deck") or {}
    s2 = deck2.get("slides") or []
    print("build_rich_assets 后：%d 页，首页 type=%s，首末标题：%s | %s"
          % (len(s2), s2[0].get("type"), s2[0].get("title"), s2[-1].get("title")))

    # 3) 真出 pptx（顺带验末页声明）
    path = "_probe_417_deck.pptx"
    deck2["notice"] = "本材料由智能助手辅助起草，已按官方要求排版，请以团队真实情况核对补充后提交。"
    build_deck(deck2, path, {})
    _last = list(Presentation(path).slides)[-1]
    _lt = "\n".join(sh.text_frame.text for sh in _last.shapes if sh.has_text_frame)
    print("末页含声明：", "智能助手辅助起草" in _lt)
    prs = Presentation(path)
    n = len(prs.slides)
    print("实际 PPTX 页数：", n)

    # 4) 一致性校验
    r = check_ppt_matches_outline(path, OUTLINE)
    print("校验 ok=%s errors=%s" % (r.get("ok"), r.get("errors")))
    print("warnings 数：%d，样例：%s" % (len(r.get("warnings") or []), (r.get("warnings") or [])[:2]))

    print("\n结论：大纲 %d 页 / PPT %d 页 -> %s"
          % (len(PAGES), n, "一致 ✅" if n == len(PAGES) else "不一致 ❌"))


if __name__ == "__main__":
    main()
