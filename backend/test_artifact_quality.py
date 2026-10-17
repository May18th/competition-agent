# -*- coding: utf-8 -*-
"""Cross-format artifact quality regression test.

Run with the bundled WorkBuddy Python. No Flask startup or network access.
"""
import json
import os
import sys
import tempfile

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "backend"))

from quality_gate import audit_artifacts

FAILED = []


def check(name, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    print("[%s] %s%s" % (status, name, (" -> " + detail) if detail else ""))
    if not condition:
        FAILED.append(name)


def _make_blank_pdf(path):
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    with open(path, "wb") as f:
        writer.write(f)


def _make_blank_pptx(path):
    from pptx import Presentation
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(path)


def main():
    report = audit_artifacts(
        os.path.join(ROOT, "artifacts", "regression.docx"),
        os.path.join(ROOT, "artifacts", "regression.pdf"),
        os.path.join(ROOT, "artifacts", "regression.pptx"),
        result={"proposal": "project data 2024 100", "rich_charts": [],
                "rich_tables": []},
        project_name="project name intentionally absent",
        mode="deep",
        base_dir=ROOT,
    )
    check("normal artifacts have no hard errors", report["ok"],
          json.dumps(report["errors"], ensure_ascii=False))
    check("all artifact metrics returned",
          all(k in report["metrics"] for k in ("docx", "pdf", "pptx")))
    check("PDF blank pages measured", report["metrics"]["pdf"]["blank_pages"] == [])
    check("PPT density measured", report["metrics"]["pptx"]["slides"] >= 8 and
          report["metrics"]["pptx"]["characters"] > 0)

    with tempfile.TemporaryDirectory(prefix="artifact_quality_") as tmp:
        pdf = os.path.join(tmp, "blank.pdf")
        pptx = os.path.join(tmp, "blank.pptx")
        _make_blank_pdf(pdf)
        _make_blank_pptx(pptx)
        broken = audit_artifacts(pdf_path=pdf, pptx_path=pptx, mode="fast")
        joined = ";".join(broken["errors"])
        check("blank PDF rejected", "PDF 存在空白页" in joined, joined)
        check("blank PPT rejected", "PPT 页数过少" in joined and "PPT 存在空白页" in joined, joined)

    if FAILED:
        print("FAIL: %s" % FAILED)
        return 1
    print("PASS: artifact quality regression")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
