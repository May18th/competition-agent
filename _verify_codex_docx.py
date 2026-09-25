# -*- coding: utf-8 -*-
"""Codex 接线验收：打真实 HTTP 接口，验证导出的 docx 是否真的达标。

用法：python _verify_codex_docx.py
只验「格式层是否接通」，不消耗模型额度（不跑生成）。
"""
import io
import json
import os
import re
import sys
import urllib.request
import uuid
from docx import Document

BASE = 'http://127.0.0.1:8080'
OUT = '_verify_out'
os.makedirs(OUT, exist_ok=True)

# 关掉系统代理，否则 127.0.0.1 会被送去代理拿到假 502
_opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))

PASS, FAIL = [], []


W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def cn_font(run):
    """取中文字体（w:rFonts 的 eastAsia）。
    注意：run.font.name 读的是 w:ascii（西文，规范要求 Times New Roman），
    直接拿它断言"中文字体"会误判。"""
    rpr = run._element.rPr
    if rpr is None:
        return None
    rf = rpr.find(W + 'rFonts')
    return rf.get(W + 'eastAsia') if rf is not None else None


def check(name, cond, detail=''):
    (PASS if cond else FAIL).append(name)
    print('  %s  %s%s' % ('[OK]  ' if cond else '[FAIL]', name, ('  -> ' + detail) if detail else ''))


def post_json(path, payload):
    req = urllib.request.Request(
        BASE + path, data=json.dumps(payload).encode('utf-8'),
        headers={'Content-Type': 'application/json'}, method='POST')
    with _opener.open(req, timeout=60) as r:
        return r.read()


def post_multipart(path, fields, file_field=None, file_path=None):
    boundary = '----wb' + uuid.uuid4().hex
    body = io.BytesIO()
    for k, v in fields.items():
        body.write(('--%s\r\nContent-Disposition: form-data; name="%s"\r\n\r\n%s\r\n'
                    % (boundary, k, v)).encode('utf-8'))
    if file_path:
        fname = os.path.basename(file_path)
        with open(file_path, 'rb') as f:
            data = f.read()
        body.write(('--%s\r\nContent-Disposition: form-data; name="%s"; filename="%s"\r\n'
                    'Content-Type: application/vnd.openxmlformats-officedocument.'
                    'wordprocessingml.document\r\n\r\n' % (boundary, file_field, fname)).encode('utf-8'))
        body.write(data)
        body.write(('\r\n--%s--\r\n' % boundary).encode('utf-8'))
    else:
        body.write(('--%s--\r\n' % boundary).encode('utf-8'))
    req = urllib.request.Request(
        BASE + path, data=body.getvalue(),
        headers={'Content-Type': 'multipart/form-data; boundary=' + boundary}, method='POST')
    with _opener.open(req, timeout=60) as r:
        return r.read()


MD = """# 智适应助老陪护机器人

## 一、项目背景

本项目面向老龄化社会的陪护需求，结合多模态感知与边缘计算技术，构建一套智适应助老陪护机器人系统，
用于解决独居老人跌倒发现滞后、用药提醒缺失、情感陪伴不足三类痛点。

## 二、技术方案

### （一）多模态感知层

采用视觉与语音融合的感知方案，在端侧完成推理，保证隐私不出户。

| 模块 | 指标 | 说明 |
|---|---|---|
| 视觉跌倒检测 | 30fps | 端侧推理，延迟低于 50ms |
| 语音唤醒 | 95% | 支持方言识别 |

## 三、总结与展望

后续将在三个社区开展试点，验证方案的可复制性与商业可行性。
"""

# ============ 场景1：申报书档位（default）============
print('\n===== 场景1：/api/export_word  doc_type=default（申报书）=====')
p1 = os.path.join(OUT, 's1_default.docx')
open(p1, 'wb').write(post_json('/api/export_word', {
    'title': '智适应助老陪护机器人',
    'text': MD, 'doc_type': 'default',
    'school': '某某大学', 'team': '启明创新团队', 'advisor': '张三 教授',
}))
d = Document(p1)
xml = d.element.body.xml
texts = [p.text for p in d.paragraphs]
check('表格是真表格（不是 | xxx | 裸文本）', len(d.tables) >= 1,
      '表格数=%d，首个 %dx%d' % (len(d.tables), len(d.tables[0].rows), len(d.tables[0].columns)))
check('无竖线残留', not any(t.strip().startswith('|') for t in texts))
check('封面含学校名', any('某某大学' in t for t in texts))
check('封面含团队名', any('启明创新团队' in t for t in texts))
check('封面含指导老师', any('张三 教授' in t for t in texts))
check('目录用了 PAGEREF 域', 'PAGEREF' in xml and 'bookmarkStart' in xml,
      'PAGEREF=%d 书签=%d' % (xml.count('PAGEREF'), xml.count('bookmarkStart')))
h1 = [p for p in d.paragraphs if p.style.name == 'Heading 1']
if h1:
    r = h1[0].runs[0]
    sz = r.font.size.pt if r.font.size else None
    check('一级标题 中文黑体 + 二号(22pt)', sz == 22 and cn_font(r) == '黑体',
          '%spt 中文=%s 西文=%s' % (sz, cn_font(r), r.font.name))
body_ps = [p for p in d.paragraphs if p.style.name == 'Normal' and len(p.text) > 30]
if body_ps:
    r = body_ps[0].runs[0]
    sz = r.font.size.pt if r.font.size else None
    check('正文 中文宋体 + 小四(12pt)', sz == 12 and cn_font(r) == '宋体',
          '%spt 中文=%s 西文=%s' % (sz, cn_font(r), r.font.name))
    check('正文首行缩进 2 字符', body_ps[0]._p.xml.find('firstLineChars="200"') > 0)
sec = d.sections[0]
check('页边距 上下2.54 左右2.5cm',
      abs(sec.top_margin.cm - 2.54) < 0.02 and abs(sec.left_margin.cm - 2.5) < 0.02,
      'T=%.2f L=%.2f' % (sec.top_margin.cm, sec.left_margin.cm))
check('页脚有页码域', 'PAGE' in sec.footer.paragraphs[0]._p.xml if sec.footer.paragraphs else False)

# ============ 场景2：大纲档位（outline）============
print('\n===== 场景2：/api/export_word  doc_type=outline（大纲）=====')
p2 = os.path.join(OUT, 's2_outline.docx')
open(p2, 'wb').write(post_json('/api/export_word', {
    'title': '路演PPT大纲', 'text': MD, 'doc_type': 'outline'}))
d2 = Document(p2)
b2 = [p for p in d2.paragraphs if p.style.name == 'Normal' and len(p.text) > 30]
check('大纲无首行缩进', b2 and b2[0]._p.xml.find('firstLineChars') < 0,
      '缩进=%s' % (b2[0].paragraph_format.first_line_indent if b2 else 'N/A'))
check('大纲无封面（不出现学校行）', not any('某某大学' in p.text for p in d2.paragraphs))
h = [p for p in d2.paragraphs if p.style.name == 'Heading 1']
if h:
    sz = h[0].runs[0].font.size.pt if h[0].runs and h[0].runs[0].font.size else None
    check('大纲一级标题 四号(14pt)', sz == 14, '%spt' % sz)

# ============ 场景3：B 链路 keep（保留原格式只换字）============
print('\n===== 场景3：/api/export_original_format  mode=keep（保留原格式）=====')
SRC = 'iCAN参赛项目计划书-护途CareWay(1).docx'
if not os.path.exists(SRC):
    check('找到样例原文件', False, SRC + ' 不存在')
else:
    src = Document(SRC)
    src_p, src_t = len(src.paragraphs), len(src.tables)
    src_cell = len(src.tables[0].rows) * len(src.tables[0].columns) if src_t else 0
    # 抓原文件首个非空段落的字体字号，用来比对是否被冲掉
    fmt_before = None
    for p in src.paragraphs:
        if p.text.strip() and p.runs:
            r = p.runs[0]
            fmt_before = (cn_font(r), r.font.size.pt if r.font.size else None, r.bold)
            break
    OPT = '\n'.join(['【已优化】第%d条：内容已由 AI 实质性改写，表述更专业、论证更充分。' % i
                     for i in range(1, 25)])
    p3 = os.path.join(OUT, 's3_keep.docx')
    open(p3, 'wb').write(post_multipart(
        '/api/export_original_format',
        {'optimized_text': OPT, 'mode': 'keep', 'title': '项目申报书',
         'school': '某某大学', 'team': '启明创新团队', 'advisor': '张三 教授'},
        file_field='original_file', file_path=SRC))
    out = Document(p3)
    check('段落数与原文件一致', len(out.paragraphs) == src_p,
          '原=%d 新=%d' % (src_p, len(out.paragraphs)))
    check('表格数与原文件一致', len(out.tables) == src_t,
          '原=%d 新=%d' % (src_t, len(out.tables)))
    if src_t and out.tables:
        check('表格行列数未变',
              len(out.tables[0].rows) * len(out.tables[0].columns) == src_cell,
              '原=%d 新=%d' % (src_cell, len(out.tables[0].rows) * len(out.tables[0].columns)))
    alltext = '\n'.join([p.text for p in out.paragraphs] +
                        [c.text for t in out.tables for r in t.rows for c in r.cells])
    check('内容被替换（出现优化后的文字）', '【已优化】第1条' in alltext)
    fmt_after = None
    for p in out.paragraphs:
        if p.text.strip() and p.runs:
            r = p.runs[0]
            fmt_after = (cn_font(r), r.font.size.pt if r.font.size else None, r.bold)
            break
    check('中文字体/字号/加粗 未被冲掉', fmt_before == fmt_after,
          '原=%s 新=%s' % (fmt_before, fmt_after))
    check('keep 模式不套封面（无目录域）', 'PAGEREF' not in out.element.body.xml)

    # ============ 场景4：B 链路 reformat ============
    print('\n===== 场景4：/api/export_original_format  mode=reformat（标准重排）=====')
    p4 = os.path.join(OUT, 's4_reformat.docx')
    open(p4, 'wb').write(post_multipart(
        '/api/export_original_format',
        # 用带标题的 Markdown：目录按标题生成，没有标题自然不会有目录
        {'optimized_text': MD, 'mode': 'reformat', 'title': '项目申报书',
         'school': '某某大学', 'team': '启明创新团队', 'advisor': '张三 教授'},
        file_field='original_file', file_path=SRC))
    o4 = Document(p4)
    t4 = [p.text for p in o4.paragraphs]
    check('reformat 有封面学校名', any('某某大学' in t for t in t4))
    check('reformat 有目录域', 'PAGEREF' in o4.element.body.xml)

print('\n================ 结果 ================')
print('通过 %d / 失败 %d' % (len(PASS), len(FAIL)))
if FAIL:
    print('失败项：')
    for f in FAIL:
        print('  - ' + f)
print('产物目录：%s' % os.path.abspath(OUT))
sys.exit(1 if FAIL else 0)
