"""Deterministic export recovery helpers. Never invent project facts."""
import re


def sanitize_export_text(text):
    """Remove unresolved author-facing placeholder fragments without fabricating values."""
    value = str(text or "")
    value = re.sub(r"【[^】]*(?:待你填写|待填写|待补充|待填|TBD|TODO)[^】]*】", "", value, flags=re.I)
    value = re.sub(r"\[(?:[^\]]*(?:待你填写|待填写|待补充|待填|TBD|TODO)[^\]]*)\]", "", value, flags=re.I)
    value = re.sub(r"[（(]\s*(?:待你填写|待填写|待补充|待填|TBD|TODO)\s*[)）]", "", value, flags=re.I)
    # These sentences are instructions to the author, not content suitable for a judge-facing export.
    value = re.sub(r"[^。！？\n]*(?:待你填写|待填写|待补充|待填|TBD|TODO)[^。！？\n]*[。！？]?", "", value, flags=re.I)
    value = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", value)
    value = re.sub(r"~~([^~\n]+)~~", r"\1", value)
    value = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"\1", value)
    value = re.sub(r"(?<!\w)_([^_\n]+)_(?!\w)", r"\1", value)
    value = re.sub(r"(?m)^[ \t]*(?:团队名称|项目名称|学校|院校|指导老师|指导教师|负责人|队长|姓名|电话|联系方式)[ \t]*[：:]\s*$", "", value)
    value = re.sub(r"(?m)^[ \t]*[【\[（(]?[ \t]*[】\]）)]?[ \t]*$", "", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _inline_plain(text):
    value = str(text or "").strip()
    value = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", value)
    value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", value)
    value = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), value)
    value = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)|(?<!_)_([^_]+)_(?!_)",
                   lambda m: m.group(1) or m.group(2), value)
    value = re.sub(r"`([^`]+)`", r"\1", value)
    value = re.sub(r"~~([^~]+)~~", r"\1", value)
    return value.strip()


def _table_rows(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [c.strip() for c in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            rows.append([_inline_plain(c) for c in cells])
        index += 1
    return rows, index


def render_recovery_docx(title, text, out_path):
    """Write a readable editable DOCX if the styled renderer cannot complete."""
    from docx import Document

    doc = Document()
    doc.add_heading(str(title or "项目材料"), 0)
    lines = str(text or "").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = _table_rows(lines, index)
            if rows:
                table = doc.add_table(rows=1, cols=max(len(row) for row in rows))
                table.style = "Table Grid"
                for row_index, row in enumerate(rows):
                    cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
                    for col_index, cell in enumerate(cells):
                        cell.text = row[col_index] if col_index < len(row) else ""
            continue
        match = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if match:
            doc.add_heading(_inline_plain(match.group(2)), min(len(match.group(1)), 3))
        elif re.fullmatch(r"(?:-{3,}|\*{3,}|_{3,})", stripped):
            pass
        else:
            stripped = re.sub(r"^>\s*", "", stripped)
            stripped = re.sub(r"^(?:[-*•]\s+|\d+[.)、]\s+)", "", stripped)
            if stripped.startswith("```"):
                index += 1
                continue
            doc.add_paragraph(_inline_plain(stripped))
        index += 1
    doc.save(out_path)
    return out_path


def render_recovery_pdf(title, text, out_path):
    """Write a plain but readable PDF if rich Markdown layout fails."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from pdf_render import _FONT, _escape_html

    title_style = ParagraphStyle("recovery-title", fontName=_FONT, fontSize=18,
                                 leading=25, alignment=TA_CENTER, spaceAfter=18)
    body_style = ParagraphStyle("recovery-body", fontName=_FONT, fontSize=10.5,
                                leading=16, spaceAfter=6, wordWrap="CJK")
    heading_style = ParagraphStyle("recovery-heading", fontName=_FONT, fontSize=13,
                                   leading=19, spaceBefore=9, spaceAfter=5, wordWrap="CJK")
    flow = [Paragraph(_escape_html(str(title or "项目材料")), title_style)]
    lines = str(text or "").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = _table_rows(lines, index)
            if rows:
                table_data = [[Paragraph(_inline_plain(cell), body_style) for cell in row] for row in rows]
                table = Table(table_data, repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA3B2")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF3")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ]))
                flow.append(table)
                flow.append(Spacer(1, 6))
            continue
        match = re.match(r"^#{1,3}\s+(.+)$", stripped)
        if match:
            flow.append(Paragraph(_escape_html(_inline_plain(match.group(1))), heading_style))
        elif re.fullmatch(r"(?:-{3,}|\*{3,}|_{3,})", stripped):
            pass
        else:
            stripped = re.sub(r"^>\s*", "", stripped)
            stripped = re.sub(r"^(?:[-*•]\s+|\d+[.)、]\s+)", "", stripped)
            if stripped.startswith("```"):
                index += 1
                continue
            flow.append(Paragraph(_escape_html(_inline_plain(stripped)), body_style))
            flow.append(Spacer(1, 2))
        index += 1
    SimpleDocTemplate(out_path, pagesize=A4, title=str(title or "项目材料")).build(flow)
    return out_path
