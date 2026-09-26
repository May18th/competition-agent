# -*- coding: utf-8 -*-
"""
验证 docx_render.py 的四档文档格式是否符合规范。
跑法：python _test_docx_render.py
"""
import os
import sys

from docx import Document
from docx.oxml.ns import qn
from docx_render import build_docx

MD = """# 一、项目背景

本项目面向老龄化社会的陪护需求,结合多模态感知与边缘计算技术(Edge AI),
构建一套智适应助老陪护机器人系统.

## (一)核心评分维度

| 评分维度 | 占比 | 核心考察点 |
|---------|------|-----------|
| 创新性 | 30% | 原创性、新颖性 |
| 实用性 | 25% | 是否解决真实问题 |
| 技术难度 | 20% | 技术复杂度与壁垒 |

表 5-1 评分权重分布表

## (二)技术路线

1. 多模态数据采集层
2. 边缘推理层
3. 自适应决策层

- 低延迟:端侧推理 < 50ms
- 高可靠:离线可用

> 注:以上为通用创新创业竞赛常识推断.

### 1. 创新点

**原创算法**是本项目的核心壁垒.

```python
def infer(x):
    return model(x)
```
"""


def _info(doc, name):
    s = doc.sections[0]
    print('\n===== %s =====' % name)
    print('页面 %.1f x %.1f cm | 边距 T%.2f B%.2f L%.2f R%.2f'
          % (s.page_width.cm, s.page_height.cm, s.top_margin.cm,
             s.bottom_margin.cm, s.left_margin.cm, s.right_margin.cm))
    print('表格数: %d  段落数: %d' % (len(doc.tables), len(doc.paragraphs)))
    if doc.tables:
        t = doc.tables[0]
        print('首个表格: %d行 x %d列, 表头=%s'
              % (len(t.rows), len(t.columns),
                 [c.text for c in t.rows[0].cells]))
    for i, p in enumerate(doc.paragraphs[:200]):
        txt = p.text.strip()
        if not txt or len(txt) > 44:
            continue
        runs = [r for r in p.runs if r.text.strip()]
        if not runs:
            continue
        r = runs[0]
        rF = r._element.rPr.rFonts if r._element.rPr is not None else None
        ea = rF.get(qn('w:eastAsia')) if rF is not None else None
        asc = rF.get(qn('w:ascii')) if rF is not None else None
        sz = r.font.size.pt if r.font.size else None
        pPr = p._p.pPr
        ind_chars = None
        if pPr is not None:
            ind = pPr.find(qn('w:ind'))
            if ind is not None:
                ind_chars = ind.get(qn('w:firstLineChars'))
        print('  %-26s %5spt %-4s/%-16s 缩进=%-5s %s'
              % (txt[:26], sz, ea, asc, ind_chars,
                 '粗' if r.font.bold else '  '))
        if i > 60:
            break
    return True


def main():
    outdir = os.path.join(os.path.dirname(os.path.abspath(__file__)), '_docx_test')
    os.makedirs(outdir, exist_ok=True)

    cases = [
        ('申报书(标准格式)', 'default', '项目申报书'),
        ('报告', 'report', '项目分析报告'),
        ('分析', 'analysis', '竞品分析'),
        ('大纲', 'outline', '路演 PPT 大纲'),
    ]
    for label, dtype, title in cases:
        doc = build_docx(title, MD, subtitle='iCAN大学生创新创业大赛',
                         org='示例团队', doc_type=dtype)
        path = os.path.join(outdir, '%s.docx' % dtype)
        doc.save(path)
        _info(doc, '%s -> %s (%.0fKB)' % (label, os.path.basename(path),
                                          os.path.getsize(path) / 1024.0))

    # 断言
    errs = []
    for label, dtype, title in cases:
        p = os.path.join(outdir, '%s.docx' % dtype)
        d = Document(p)
        s = d.sections[0]
        if abs(s.page_width.cm - 21.0) > 0.1:
            errs.append('%s 页宽不是 A4' % label)
        if abs(s.top_margin.cm - 2.54) > 0.01 or abs(s.left_margin.cm - 3.17) > 0.01:
            errs.append('%s 页边距不符规范' % label)
        if not d.tables:
            errs.append('%s 表格没渲染成真表格' % label)
        else:
            for c in d.tables[0].rows[0].cells:
                if '|' in c.text:
                    errs.append('%s 表格里还残留竖线' % label)
                    break
        # 字体检查：只能宋体/黑体
        for para in d.paragraphs:
            for r in para.runs:
                rF = r._element.rPr.rFonts if r._element.rPr is not None else None
                if rF is None:
                    continue
                ea = rF.get(qn('w:eastAsia'))
                # 中文只允许宋体/黑体；Times New Roman 仅用于代码块等纯西文内容
                if ea and ea not in ('宋体', '黑体', 'Times New Roman'):
                    errs.append('%s 出现非规范中文字体: %s' % (label, ea))
                    break

    # 大纲不应首行缩进
    d = Document(os.path.join(outdir, 'outline.docx'))
    for para in d.paragraphs:
        pPr = para._p.pPr
        if pPr is None:
            continue
        ind = pPr.find(qn('w:ind'))
        if ind is not None and ind.get(qn('w:firstLineChars')) == '200':
            errs.append('大纲档位仍有首行缩进')
            break

    # 申报书/报告 正文应首行缩进 2 字符
    d = Document(os.path.join(outdir, 'report.docx'))
    has_indent = False
    for para in d.paragraphs:
        pPr = para._p.pPr
        if pPr is None:
            continue
        ind = pPr.find(qn('w:ind'))
        if ind is not None and ind.get(qn('w:firstLineChars')) == '200':
            has_indent = True
            break
    if not has_indent:
        errs.append('报告档位缺少首行缩进')

    print('\n================ 断言 ================')
    if errs:
        for e in errs:
            print('  [FAIL]', e)
        return 1
    print('  全部通过：页边距/字体/表格/缩进 均符合规范')
    print('  输出目录:', outdir)
    return 0


if __name__ == '__main__':
    sys.exit(main())
