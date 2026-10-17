import os
import tempfile
import unittest

from export_recovery import render_recovery_docx, render_recovery_pdf, sanitize_export_text
from export_quality import (
    hard_artifact_issues,
    hard_content_issues,
    public_artifact_reasons,
    verify_artifact_integrity,
)


class TestExportRecovery(unittest.TestCase):
    def test_hard_content_issues_are_separated_from_soft_artifact_warnings(self):
        self.assertIn("正文为空", hard_content_issues("" )[0])
        self.assertTrue(hard_content_issues("乱码???????"))
        self.assertTrue(hard_content_issues("四川大学团队方案", double_blind=True))
        self.assertFalse(hard_content_issues("章节缺失但正文可读"))
        self.assertEqual(hard_artifact_issues("pptx", ["PPT 页数过少（5 < 8）"]), [])
        self.assertEqual(len(hard_artifact_issues("pptx", ["pptx 检查失败：损坏"])), 1)
        self.assertEqual(
            public_artifact_reasons("docx", ["docx 检查失败：BadZipFile /var/private/path"]),
            ["Word 文件无法完整读取，系统已尝试自动修复"],
        )

    def test_sanitizer_removes_instructions_without_inventing_values(self):
        source = "## 项目概述\n真实内容保留。\n团队名称：【待你填写：团队名称】\n该数据待补充。\n"
        result = sanitize_export_text(source)
        self.assertIn("真实内容保留。", result)
        self.assertNotIn("待你填写", result)
        self.assertNotIn("待补充", result)
        self.assertNotIn("团队名称：", result)

    def test_fallback_renderers_create_readable_artifacts(self):
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = os.path.join(tmp, "recovery.docx")
            pdf_path = os.path.join(tmp, "recovery.pdf")
            text = "## 项目概述\n项目**正文**内容和[来源](https://example.invalid)。\n\n|字段|结果|\n|---|---|\n|指标|已验证|\n\n```text\n恢复数据\n```"
            render_recovery_docx("恢复稿", text, docx_path)
            render_recovery_pdf("恢复稿", text, pdf_path)
            self.assertGreater(os.path.getsize(docx_path), 0)
            self.assertGreater(os.path.getsize(pdf_path), 0)
            self.assertTrue(verify_artifact_integrity("docx", docx_path))
            self.assertTrue(verify_artifact_integrity("pdf", pdf_path))
            from docx import Document
            doc = Document(docx_path)
            doc_text = "\n".join(p.text for p in doc.paragraphs)
            doc_text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            from pypdf import PdfReader
            pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf_path).pages)
            self.assertIn("项目正文内容和来源", doc_text)
            self.assertIn("指标", doc_text)
            self.assertNotRegex(doc_text, r"#{1,6}\s|\*\*|\|字段\||```|\[[^\]]+\]\(")
            self.assertIn("项目正文内容和来源", pdf_text)
            self.assertIn("指标", pdf_text)
            self.assertNotRegex(pdf_text, r"#{1,6}\s|\*\*|\|字段\||```|\[[^\]]+\]\(")

    def test_integrity_fallback_rejects_corrupt_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            corrupt = os.path.join(tmp, "broken.docx")
            with open(corrupt, "wb") as target:
                target.write(b"not a docx")
            self.assertFalse(verify_artifact_integrity("docx", corrupt))

    def test_production_renderers_leave_no_markdown_syntax(self):
        from docx import Document
        from docx_render import build_docx
        from pdf_render import render_markdown_to_pdf
        from pypdf import PdfReader

        source = "## 项目概述\n项目**正文**内容和[来源](https://example.invalid)。\n\n|字段|结果|\n|---|---|\n|指标|已验证|\n\n```text\n恢复数据\n```"
        text = sanitize_export_text(source)
        with tempfile.TemporaryDirectory() as tmp:
            docx_path = os.path.join(tmp, "styled.docx")
            pdf_path = os.path.join(tmp, "styled.pdf")
            build_docx("测试稿", text, doc_type="analysis", cover=False, toc=False).save(docx_path)
            render_markdown_to_pdf(text, pdf_path, title="测试稿")

            doc = Document(docx_path)
            doc_text = "\n".join(p.text for p in doc.paragraphs)
            doc_text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf_path).pages)
            self.assertNotRegex(doc_text, r"#{1,6}\s|\*\*|\|字段\||```|\[[^\]]+\]\(")
            self.assertNotRegex(pdf_text, r"#{1,6}\s|\*\*|\|字段\||```|\[[^\]]+\]\(")

    def test_ppt_text_cleaner_removes_markdown_syntax(self):
        from pptx_builder import _wb_strip_md

        cleaned = _wb_strip_md("## **重点** [来源](https://example.invalid) ~~旧表述~~ _说明_")
        self.assertIn("重点", cleaned)
        self.assertIn("来源（https://example.invalid）", cleaned)
        self.assertIn("旧表述", cleaned)
        self.assertIn("说明", cleaned)
        self.assertNotRegex(cleaned, r"#{1,6}\s|\*\*|~~|\[[^\]]+\]\(|(?<!\w)_[^_]+_(?!\w)")


if __name__ == "__main__":
    unittest.main()
