import http.cookiejar
import json
import os
import tempfile
import threading
import unittest
import uuid
from urllib.error import HTTPError
from urllib.request import Request, build_opener


class ProposalRevisionHttpTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory(prefix="competition-http-e2e-")
        cls.old_cwd = os.getcwd()
        cls.old_env = {key: os.environ.get(key) for key in (
            "COMPETITION_DB_PATH", "COMPETITION_UID_SECRET_PATH", "COMPETITION_CHAT_DB_PATH",
            "COMPETITION_QUOTA_PATH", "COMPETITION_LLM_USAGE_PATH", "COMPETITION_SKIP_KB_SEED")
        }
        cls.db_path = os.path.join(cls.temp_dir.name, "isolated.sqlite")
        os.environ["COMPETITION_DB_PATH"] = cls.db_path
        os.environ["COMPETITION_UID_SECRET_PATH"] = os.path.join(cls.temp_dir.name, "uid.secret")
        os.environ["COMPETITION_CHAT_DB_PATH"] = os.path.join(cls.temp_dir.name, "chat_history.sqlite")
        os.environ["COMPETITION_QUOTA_PATH"] = os.path.join(cls.temp_dir.name, "quota.json")
        os.environ["COMPETITION_LLM_USAGE_PATH"] = os.path.join(cls.temp_dir.name, "llm_usage.sqlite")
        os.environ["COMPETITION_SKIP_KB_SEED"] = "1"
        os.chdir(cls.temp_dir.name)
        try:
            import sys
            if hasattr(sys.stdout, "reconfigure"):
                sys.stdout.reconfigure(encoding="utf-8", errors="replace")
            import app as app_module
            from werkzeug.serving import make_server

            cls.app = app_module
            import competition_agents
            cls.old_app_db_path = cls.app.DB_PATH
            cls.old_chat_db_path = cls.app._CHAT_DB_PATH
            cls.old_quota_path = competition_agents._QUOTA_FILE
            cls.old_usage_path = competition_agents._LLM_USAGE_DB_PATH
            cls.app.DB_PATH = cls.db_path
            cls.app._CHAT_DB_PATH = os.environ["COMPETITION_CHAT_DB_PATH"]
            competition_agents._QUOTA_FILE = os.environ["COMPETITION_QUOTA_PATH"]
            competition_agents._LLM_USAGE_DB_PATH = os.environ["COMPETITION_LLM_USAGE_PATH"]
            cls.app.app.config.update(TESTING=True)
            cls.app.init_db()
            cls.server = make_server("127.0.0.1", 0, cls.app.app, threaded=True)
            cls.base_url = "http://127.0.0.1:%d" % cls.server.server_port
            cls.server_thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
            cls.server_thread.start()
            cls.cookie_jar = http.cookiejar.CookieJar()
            from urllib.request import HTTPCookieProcessor
            cls.http = build_opener(HTTPCookieProcessor(cls.cookie_jar))
            cls._request("/api/me")
            signed_cookie = next(c.value for c in cls.cookie_jar if c.name == "comp_uid")
            cls.uid = cls.app._unsign_uid(signed_cookie)
            cls.chart_name = "e2e_%s.png" % uuid.uuid4().hex
            cls.chart_path = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))),
                                          "generated", cls.chart_name)
            os.makedirs(os.path.dirname(cls.chart_path), exist_ok=True)
            from PIL import Image
            Image.new("RGB", (640, 360), (42, 150, 132)).save(cls.chart_path, format="PNG")
        except Exception:
            os.chdir(cls.old_cwd)
            if hasattr(cls, "old_app_db_path"):
                cls.app.DB_PATH = cls.old_app_db_path
                cls.app._CHAT_DB_PATH = cls.old_chat_db_path
                import competition_agents
                competition_agents._QUOTA_FILE = cls.old_quota_path
                competition_agents._LLM_USAGE_DB_PATH = cls.old_usage_path
            for key, old in cls.old_env.items():
                if old is None:
                    os.environ.pop(key, None)
                else:
                    os.environ[key] = old
            if hasattr(cls, "old_app_db_path"):
                cls.app.DB_PATH = cls.old_app_db_path
                cls.app._CHAT_DB_PATH = cls.old_chat_db_path
                import competition_agents
                competition_agents._QUOTA_FILE = cls.old_quota_path
                competition_agents._LLM_USAGE_DB_PATH = cls.old_usage_path
            cls.temp_dir.cleanup()
            raise

    @classmethod
    def tearDownClass(cls):
        try:
            cls.server.shutdown()
            cls.server_thread.join(timeout=5)
        finally:
            if os.path.exists(cls.chart_path):
                os.remove(cls.chart_path)
            os.chdir(cls.old_cwd)
            for key, old in cls.old_env.items():
                if old is None:
                    os.environ.pop(key, None)
                else:
                    os.environ[key] = old
            cls.temp_dir.cleanup()

    @classmethod
    def _request(cls, path, payload=None, method=None):
        body = None if payload is None else json.dumps(payload, ensure_ascii=False).encode("utf-8")
        request = Request(cls.base_url + path, data=body, method=method or ("POST" if body else "GET"))
        if body is not None:
            request.add_header("Content-Type", "application/json")
        try:
            response = cls.http.open(request, timeout=45)
            return response.status, response.headers, response.read()
        except HTTPError as error:
            return error.code, error.headers, error.read()

    def setUp(self):
        self.long_text = "## 项目概述\n" + ("项目解决真实需求，采用可验证的实施方案并面向评委说明价值。\n" * 18)

    def test_revision_preview_apply_undo_and_owner_isolation(self):
        import competition_agents
        from types import SimpleNamespace
        from unittest.mock import patch

        history_id = self.app.save_history_item({
            "competition_name": "测试赛事", "idea": "测试项目", "mode": "fast", "user_id": self.uid,
            "data": {"proposal": self.long_text},
        })
        revised = "## 项目概述\n" + ("修改后面向评委清楚说明问题、证据、方案和可验证成效。\n" * 18)
        end_marker = self.app._REVISION_END_MARKER
        class FakeLLM:
            def invoke(self, *_args, **_kwargs):
                return SimpleNamespace(content=revised + end_marker)

        with patch.object(competition_agents, "llm", FakeLLM()):
            status, _headers, body = self._request("/api/revise_chapter", {
                "instruction": "突出评委关心的证据和成效", "history_id": history_id,
                "competition_name": "测试赛事", "mode": "fast",
            })
        result = json.loads(body)
        self.assertEqual(status, 200)
        self.assertEqual(result["old_proposal"], self.long_text)
        self.assertEqual(result["proposal"], revised.strip())
        self.assertEqual(self._history_proposal(history_id), self.long_text, "预览不能覆盖原稿")

        status, _headers, body = self._request("/api/proposal_revision/%d" % history_id, {
            "action": "apply", "proposal": revised, "competition_name": "测试赛事",
        })
        self.assertEqual(status, 200, body.decode("utf-8"))
        self.assertEqual(self._history_proposal(history_id), revised.strip())

        status, _headers, body = self._request("/api/export_word", {
            "title": "应用后的修改稿", "text": self.long_text, "proposal": self.long_text,
            "history_id": history_id, "competition_name": "测试赛事", "mode": "fast",
        })
        self.assertEqual(status, 200, body.decode("utf-8", errors="replace")[:300])
        applied_doc = os.path.join(self.temp_dir.name, "applied.docx")
        with open(applied_doc, "wb") as output:
            output.write(body)
        from docx import Document
        applied_text = "\n".join(p.text for p in Document(applied_doc).paragraphs)
        self.assertIn("修改后面向评委", applied_text)
        self.assertNotIn("项目解决真实需求", applied_text)

        status, _headers, body = self._request("/api/proposal_revision/%d" % history_id, {"action": "undo"})
        self.assertEqual(status, 200, body.decode("utf-8"))
        self.assertEqual(self._history_proposal(history_id), self.long_text)

        foreign_id = self.app.save_history_item({
            "competition_name": "测试赛事", "mode": "fast", "user_id": "another-user",
            "data": {"proposal": self.long_text},
        })
        status, _headers, _body = self._request("/api/proposal_revision/%d" % foreign_id,
                                                 {"action": "apply", "proposal": revised})
        self.assertEqual(status, 404)

    def test_long_revision_preserves_unselected_sections_and_rejects_truncation(self):
        import competition_agents
        from types import SimpleNamespace
        from unittest.mock import patch

        tail = "## 数据附录\n" + ("原始测算数据和指标口径必须保留。\n" * 500)
        proposal = "## 项目概述\n旧方案基于调研完成原型验证。\n\n" + tail
        history_id = self.app.save_history_item({
            "competition_name": "测试赛事", "mode": "fast", "user_id": self.uid,
            "data": {"proposal": proposal},
        })
        complete = SimpleNamespace(content=(
            "## 项目概述\n新版方案补充可复核的原型验证结果。" + self.app._REVISION_END_MARKER))
        with patch.object(competition_agents, "llm", SimpleNamespace(invoke=lambda *_a, **_k: complete)):
            status, _headers, body = self._request("/api/revise_chapter", {
                "history_id": history_id, "instruction": "突出可复核证据", "chapter": "项目概述",
            })
        self.assertEqual(status, 200, body.decode("utf-8", errors="replace"))
        self.assertTrue(json.loads(body)["proposal"].endswith(tail))

        incomplete = SimpleNamespace(content="## 项目概述\n截断的结果")
        with patch.object(competition_agents, "llm", SimpleNamespace(invoke=lambda *_a, **_k: incomplete)):
            status, _headers, body = self._request("/api/revise_chapter", {
                "history_id": history_id, "instruction": "优化项目概述", "chapter": "项目概述",
            })
        self.assertEqual(status, 502)
        self.assertEqual(json.loads(body)["error_code"], "incomplete_revision")
        self.assertEqual(self._history_proposal(history_id), proposal)

    def test_revision_can_target_third_level_heading(self):
        import competition_agents
        from types import SimpleNamespace
        from unittest.mock import patch

        untouched = "### 数据分析\n原有分析步骤和口径保持不变。\n"
        proposal = "## 实施方案\n### 数据采集\n原有采集方案。\n" + untouched
        history_id = self.app.save_history_item({
            "competition_name": "测试赛事", "mode": "fast", "user_id": self.uid,
            "data": {"proposal": proposal},
        })
        result = SimpleNamespace(content=(
            "### 数据采集\n更新后的采集方案。" + self.app._REVISION_END_MARKER))
        with patch.object(competition_agents, "llm", SimpleNamespace(invoke=lambda *_a, **_k: result)):
            status, _headers, body = self._request("/api/revise_chapter", {
                "history_id": history_id, "instruction": "完善采集方案", "chapter": "数据采集",
            })
        self.assertEqual(status, 200, body.decode("utf-8", errors="replace"))
        revised = json.loads(body)["proposal"]
        self.assertIn("更新后的采集方案", revised)
        self.assertTrue(revised.endswith(untouched))

    def test_apply_rejects_missing_sections_and_double_blind_identity(self):
        from unittest.mock import patch

        history_id = self.app.save_history_item({
            "competition_name": "测试赛事", "mode": "fast", "user_id": self.uid,
            "data": {"proposal": "## 项目概述\n原始方案。\n## 市场分析\n原始分析。"},
        })
        status, _headers, body = self._request("/api/proposal_revision/%d" % history_id, {
            "action": "apply", "proposal": "## 项目概述\n只剩一个章节。",
        })
        self.assertEqual(status, 422)
        self.assertIn("缺少原有章节", json.loads(body)["error"])

        with patch.object(self.app, "_is_double_blind", return_value=True):
            status, _headers, body = self._request("/api/proposal_revision/%d" % history_id, {
                "action": "apply",
                "proposal": "## 项目概述\n方案仍包含指导老师：张三。\n## 市场分析\n依据公开资料。",
            })
        self.assertEqual(status, 422)
        self.assertIn("安全校验", json.loads(body)["error"])
        self.assertEqual(self._history_proposal(history_id), "## 项目概述\n原始方案。\n## 市场分析\n原始分析。")

    def test_admin_llm_usage_endpoint_requires_admin_and_aggregates(self):
        import competition_agents
        from llm_usage_store import record_usage
        from unittest.mock import patch

        record_usage(competition_agents._LLM_USAGE_DB_PATH, user_id=self.uid,
                     stage="/api/revise_chapter", section="市场分析", mode="deep",
                     model="test-model", prompt_tokens=20, completion_tokens=8)
        status, _headers, _body = self._request("/api/admin/llm-usage")
        self.assertEqual(status, 401)
        with patch.object(self.app, "_require_admin", return_value=None):
            status, _headers, body = self._request("/api/admin/llm-usage?days=1")
        self.assertEqual(status, 200)
        result = json.loads(body)
        self.assertEqual(result["totals"]["total_tokens"], 28)
        self.assertEqual(result["sections"][0]["section"], "市场分析")

    def test_word_pdf_http_exports_cleanup_markdown_embed_chart_and_warn_soft_quality(self):
        from docx import Document
        from pypdf import PdfReader

        proposal = (
            "## 項目概述\n" + ("**面向評委**的项目说明，保留证据和方案。\n" * 10) +
            "\n团队名称：【待你填写：团队名称】\n\n"
            "## 市场分析\n" + ("市场规模和用户需求需要用公开来源验证。\n" * 10) +
            "\n| 指标 | 结果 |\n|---|---|\n| 覆盖人数 | 120 |\n"
        )
        payload = {
            "title": "测试申报书", "text": proposal, "proposal": proposal,
            "competition_name": "普通测试赛事", "mode": "fast", "use_submitted_proposal": True,
            "charts": [{"title": "市场规模趋势", "caption": "按公开数据测算",
                        "url": "/generated/" + self.chart_name, "spec": {"type": "bar"}}],
        }
        status, headers, body = self._request("/api/export_word", payload)
        self.assertEqual(status, 200, body.decode("utf-8", errors="replace")[:500])
        self.assertIn("application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                      headers.get("Content-Type", ""))
        self.assertTrue(headers.get("X-Export-Warnings"), "低分但可交付的文件应携带质量警告")
        docx_path = os.path.join(self.temp_dir.name, "http.docx")
        with open(docx_path, "wb") as output:
            output.write(body)
        document = Document(docx_path)
        document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        document_text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertNotIn("待你填写", document_text)
        self.assertNotRegex(document_text, r"#{1,6}\s|\*\*|\|指标\||\[[^\]]+\]\(")
        self.assertGreaterEqual(len(document.inline_shapes), 1, "Word 导出应真正嵌入图表图片")

        status, headers, body = self._request("/api/export_pdf", payload)
        self.assertEqual(status, 200, body.decode("utf-8", errors="replace")[:500])
        self.assertIn("application/pdf", headers.get("Content-Type", ""))
        pdf_path = os.path.join(self.temp_dir.name, "http.pdf")
        with open(pdf_path, "wb") as output:
            output.write(body)
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf_path).pages)
        self.assertNotIn("待你填写", pdf_text)
        self.assertNotRegex(pdf_text, r"#{1,6}\s|\*\*|\|指标\||\[[^\]]+\]\(")

    def test_double_blind_identity_is_removed_before_word_delivery(self):
        payload = {"title": "双盲申报书", "text": "## 项目概述\n" + self.long_text + "\n指导老师：张三",
                   "competition_name": "iCAN大学生创新创业大赛", "mode": "fast",
                   "use_submitted_proposal": True}
        from unittest.mock import patch
        with patch.object(self.app, "_is_double_blind", return_value=True):
            status, _headers, body = self._request("/api/export_word", payload)
        self.assertEqual(status, 200, body.decode("utf-8", errors="replace")[:500])
        output_path = os.path.join(self.temp_dir.name, "blind.docx")
        with open(output_path, "wb") as output:
            output.write(body)
        from docx import Document
        content = "\n".join(p.text for p in Document(output_path).paragraphs)
        self.assertNotIn("张三", content)

    def _history_proposal(self, history_id):
        conn = self.app.get_db()
        row = conn.execute("SELECT result_data FROM history WHERE id=?", (history_id,)).fetchone()
        conn.close()
        return json.loads(row[0])["proposal"]


if __name__ == "__main__":
    unittest.main()
