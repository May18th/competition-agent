# -*- coding: utf-8 -*-
"""探针：确认线上服务（8080）加载的是**新版** pptx_builder。

原理：给一张含 12 条要点的 slide，新版会自动拆成 2 页（8 + 4），
旧版只有 1 页且第 9 条之后被丢弃。看返回 pptx 的页数即可判定。
"""
import json, os, urllib.request
from pptx import Presentation

BASE = os.path.dirname(os.path.abspath(__file__))
DECK = {
    "project": "探针", "one_liner": "probe", "competition": "probe",
    "slides": [
        {"type": "cover", "title": "探针"},
        {"type": "bullets", "title": "容量探针",
         "bullets": ["要点%02d：这是一条用于测试自动分页的要点内容" % i for i in range(1, 13)]},
        {"type": "table", "title": "表格探针",
         "table": {"header": ["列A", "列B"],
                   "rows": [["行%d" % i, "值%d" % i] for i in range(1, 12)]}},
        {"type": "closing", "title": "完"},
    ],
}
payload = {"deck": DECK, "charts": [], "tables": [], "idea": "probe", "one_liner": "probe"}
req = urllib.request.Request(
    "http://127.0.0.1:8080/api/export_pptx",
    data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
    headers={"Content-Type": "application/json"}, method="POST")
opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
data = opener.open(req, timeout=60).read()
p = os.path.join(BASE, "generated", "_probe.pptx")
open(p, "wb").write(data)
prs = Presentation(p)
print("HTTP %d B, 页数 = %d" % (len(data), len(prs.slides)))
full = []
for s in prs.slides:
    for sh in s.shapes:
        if sh.has_text_frame:
            full.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    full.append(c.text)
txt = "\n".join(full)
for k in ["要点01", "要点08", "要点09", "要点12", "行1", "行11"]:
    print("  含 %-6s: %s" % (k, k in txt))
print("\n判定：", end="")
if "要点12" in txt and len(prs.slides) >= 5:
    print("✅ 新版（自动分页生效，12 条要点全部保留）")
else:
    print("⚠️  疑似旧版模块被缓存，需要重启服务")
