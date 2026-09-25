# -*- coding: utf-8 -*-
"""
docx_render.py —— 把 Markdown 渲染成符合《全国大学生竞赛申报书统一字体格式标准》的 Word 文档。

执行的标准（用户 2026-09-25 提供，优先级最高）：
  页面   A4；上下 2.54cm、左右 2.5cm；全文 1.5 倍行距
  封面   主标题黑体小初加粗居中 / 副标题黑体二号居中 / 日期等黑体小三居中；无页码
  一级   黑体二号加粗左对齐无缩进      （一、二、三、）
  二级   黑体小三加粗左对齐            （（一）（二））
  三级   黑体四号加粗左对齐            （1. 2. 3.）
  四级   宋体小四加粗                  （（1）（2））
  正文   宋体小四，1.5 倍行距，首行缩进 2 字符，段前段后 0
  西文   数字/英文统一 Times New Roman
  图表   表题/图题宋体五号加粗居中；注释宋体五号；表内文字宋体小四
  红线   无底色、无阴影、无特效；不自动编号；图片不拉伸；禁通篇加粗；中文全角标点

早期版本走过的弯路（保留在此以警醒）：
  曾按「好看」优先做成微软雅黑 + 蓝底表头 + 斑马纹 + 引用灰底 + 封面背景图，
  全部违反上面的红线。现在默认走合规路线，花哨效果一律默认关闭。

对外主入口：
    from docx_render import build_docx
    doc = build_docx('项目申报书', md_text, subtitle='iCAN大学生创新创业大赛')
    doc.save('申报书.docx')
"""
import datetime
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---- 字体：只允许宋体 / 黑体，西文统一 Times New Roman ----
CN_SONG = '宋体'
CN_HEI = '黑体'
EN = 'Times New Roman'

# ---- 中文字号 → 磅 ----
SZ_XIAOCHU = 36      # 小初
SZ_ER = 22           # 二号
SZ_XIAOSAN = 15      # 小三
SZ_SI = 14           # 四号
SZ_XIAOSI = 12       # 小四
SZ_WU = 10.5         # 五号

BLACK = RGBColor(0, 0, 0)

# ---- 文档类型样式档位（用户 2026-09-25 补充）----
# outline  大纲：最轻最简。一级黑体四号加粗 / 二级宋体小四加粗 / 三级宋体小四常规，
#                不要首行缩进、靠左对齐，不要封面和目录
# analysis 分析：标准一二三级 + 正文小四宋体 1.5 倍 + 首行缩进
# report   报告：最正式，标准一二三级 + 正文小四宋体
# default  申报书：同标准层级
_STD_HEADINGS = {
    1: (SZ_ER, CN_HEI, True),        # 黑体二号加粗
    2: (SZ_XIAOSAN, CN_HEI, True),   # 黑体小三加粗
    3: (SZ_SI, CN_HEI, True),        # 黑体四号加粗
    4: (SZ_XIAOSI, CN_SONG, True),   # 宋体小四加粗
}
PROFILES = {
    'outline': {
        'headings': {
            1: (SZ_SI, CN_HEI, True),          # 黑体四号加粗
            2: (SZ_XIAOSI, CN_SONG, True),     # 宋体小四加粗
            3: (SZ_XIAOSI, CN_SONG, False),    # 宋体小四常规
            4: (SZ_XIAOSI, CN_SONG, False),
        },
        'body_indent': False,                   # 不首行缩进，靠左对齐
        'body_align': WD_ALIGN_PARAGRAPH.LEFT,
        'cover': False, 'toc': False,
    },
    'analysis': {'headings': _STD_HEADINGS, 'body_indent': True,
                 'body_align': WD_ALIGN_PARAGRAPH.JUSTIFY,
                 'cover': True, 'toc': False},
    'report': {'headings': _STD_HEADINGS, 'body_indent': True,
               'body_align': WD_ALIGN_PARAGRAPH.JUSTIFY,
               'cover': True, 'toc': True},
    'default': {'headings': _STD_HEADINGS, 'body_indent': True,
                'body_align': WD_ALIGN_PARAGRAPH.JUSTIFY,
                'cover': True, 'toc': True},
}


# ============ 底层工具 ============

def _font(run, size, cn=CN_SONG, bold=False, color=None):
    """设置字体：西文 Times New Roman + 中文宋体/黑体。
    规范：数字、英文统一 Times New Roman；中文只允许宋体/黑体。"""
    run.font.name = EN                 # w:ascii / w:hAnsi
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), EN)
    rFonts.set(qn('w:hAnsi'), EN)
    rFonts.set(qn('w:eastAsia'), cn)
    return run


def _indent_chars(p, chars=200):
    """首行缩进 N 个字符（按字符单位，任何字号都是精确 2 字符）。"""
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), str(chars))


def _hanging_indent(p, chars=200):
    """悬挂缩进：首行顶格，换行后缩进 N 字符（列表项用）。"""
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), '-%d' % chars)
    ind.set(qn('w:leftChars'), str(chars))


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def _set_cell_margins(table, top=40, bottom=40, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for tag, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        node = OxmlElement('w:' + tag)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        mar.append(node)
    tblPr.append(mar)


# ============ 全角标点 ============

def _is_cn(ch):
    return '\u4e00' <= ch <= '\u9fff'


_PUNCT_MAP = {',': '，', ';': '；', ':': '：', '!': '！', '?': '？'}


def _fullwidth(text):
    """把中文语境里的半角标点转成全角（英文句子、小数、URL 不动）。"""
    if not text:
        return text
    out = []
    n = len(text)
    for i, ch in enumerate(text):
        prev = text[i - 1] if i > 0 else ''
        nxt = text[i + 1] if i + 1 < n else ''
        cn_near = _is_cn(prev) or _is_cn(nxt)
        if ch in _PUNCT_MAP and cn_near:
            out.append(_PUNCT_MAP[ch])
        elif ch == '.' and _is_cn(prev) and not nxt.isdigit():
            out.append('。')
        elif ch == '.' and _is_cn(nxt):
            out.append('。')
        elif ch == '(' and _is_cn(nxt):
            out.append('（')
        elif ch == ')' and _is_cn(prev):
            out.append('）')
        else:
            out.append(ch)
    return ''.join(out)


# ============ Markdown 块解析 ============

_RE_CODE = re.compile(r'^```')
_RE_HR = re.compile(r'^(-{3,}|\*{3,}|_{3,})$')
_RE_HEAD = re.compile(r'^(#{1,6})\s+(.*)$')
_RE_OL = re.compile(r'^(\d+)[.、)]\s+(.*)$')
_RE_UL = re.compile(r'^[-*•·＋]\s+(.*)$')
_RE_SEP = re.compile(r'^\|[\s:\-|]+\|$')
_RE_CAPTION = re.compile(r'^[表图]\s*[\d\-—.]')


def _is_sep(line):
    return bool(_RE_SEP.match(line)) and '-' in line


def _split_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


def _is_block_start(line):
    if not line:
        return True
    return bool(_RE_CODE.match(line) or _RE_HR.match(line) or _RE_HEAD.match(line)
                or _RE_OL.match(line) or _RE_UL.match(line)
                or line.startswith(('>', '|')))


def _blocks(text):
    """切块：('h',level,text) / ('p',[lines]) / ('ul',[items]) / ('ol',[items])
       / ('quote',[lines]) / ('code',[lines],lang) / ('table',header,rows) / ('hr',)"""
    lines = text.replace('\r\n', '\n').split('\n')
    out, i, n = [], 0, len(lines)
    while i < n:
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if _RE_CODE.match(s):
            i += 1
            buf = []
            while i < n and not _RE_CODE.match(lines[i].strip()):
                buf.append(lines[i].rstrip())
                i += 1
            i += 1
            out.append(('code', buf))
            continue
        if _RE_HR.match(s):
            out.append(('hr',))
            i += 1
            continue
        m = _RE_HEAD.match(s)
        if m:
            out.append(('h', min(len(m.group(1)), 4), m.group(2).strip()))
            i += 1
            continue
        if s.startswith('|') and i + 1 < n and _is_sep(lines[i + 1].strip()):
            header = _split_row(s)
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                rows.append(_split_row(lines[i].strip()))
                i += 1
            out.append(('table', header, rows))
            continue
        if s.startswith('>'):
            buf = []
            while i < n and lines[i].strip().startswith('>'):
                buf.append(lines[i].strip().lstrip('>').strip())
                i += 1
            out.append(('quote', buf))
            continue
        if _RE_OL.match(s):
            buf = []
            while i < n:
                mm = _RE_OL.match(lines[i].strip())
                if not mm:
                    break
                buf.append(mm.group(2).strip())
                i += 1
            out.append(('ol', buf))
            continue
        if _RE_UL.match(s):
            buf = []
            while i < n:
                mm = _RE_UL.match(lines[i].strip())
                if not mm:
                    break
                buf.append(mm.group(1).strip())       # _RE_UL 只有 1 个捕获组
                i += 1
            out.append(('ul', buf))
            continue
        buf = []
        while i < n:
            t = lines[i].strip()
            if not t or _is_block_start(t):
                break
            buf.append(t)
            i += 1
        if buf:
            out.append(('p', buf))
    return out


def _plain(text):
    """去掉行内标记（表格单元格用）。"""
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    return text.strip()


# ============ 页面 ============

def _setup_heading_styles(doc, prof):
    """把 Word 的 Heading 1~4 样式改成符合规范的字体字号。

    关键：目录（TOC 域）只认「标题样式」的段落，认不出普通段落加粗。
    所以标题必须挂上 Heading 样式，再改样式本身的外观，
    这样目录才能自动对应到正文并生成正确页码。
    """
    for level in (1, 2, 3, 4):
        try:
            st = doc.styles['Heading %d' % level]
        except KeyError:
            continue
        size, cn, bold = prof['headings'][level]
        st.font.name = EN
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = BLACK          # 去掉 Word 标题默认的主题蓝
        st.font.italic = False
        st.font.underline = False
        rPr = st.element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn('w:ascii'), EN)
        rFonts.set(qn('w:hAnsi'), EN)
        rFonts.set(qn('w:eastAsia'), cn)
        pf = st.paragraph_format
        pf.space_before = Pt(14 if level == 1 else 10)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.5
        pf.keep_with_next = True
        pf.left_indent = Cm(0)
        pf.first_line_indent = Cm(0)


def _setup_page(doc, prof=None):
    """A4 + 规范页边距 + 默认正文样式 + 标题样式（供目录域抓取）；封面页无页码。"""
    prof = prof or PROFILES['default']
    doc.sections[0].different_first_page_header_footer = True

    style = doc.styles['Normal']
    style.font.name = EN
    style.font.size = Pt(SZ_XIAOSI)
    rFonts = style.element.rPr.rFonts
    rFonts.set(qn('w:eastAsia'), CN_SONG)
    rFonts.set(qn('w:ascii'), EN)
    rFonts.set(qn('w:hAnsi'), EN)
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(2.54)
        sec.bottom_margin = Cm(2.54)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
        sec.header_distance = Cm(1.5)
        sec.footer_distance = Cm(1.5)

    _setup_heading_styles(doc, prof)


def _page_number_footer(doc):
    """页码：底部居中（封面因 different_first_page 自动无页码）。"""
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run('— '), SZ_WU, CN_SONG)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE  \\* MERGEFORMAT')
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(SZ_WU * 2)))
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = '1'
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    _font(p.add_run(' —'), SZ_WU, CN_SONG)


# ============ 封面 ============

def _header(doc, title):
    """页眉：文档标题，宋体五号居中（无装饰线，符合「无特效」要求）。"""
    header = doc.sections[0].header
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    _font(p.add_run(title or ''), SZ_WU, CN_SONG)


def _cover(doc, title, subtitle=None, org=None, school=None, team=None, advisor=None):
    """封面：主标题黑体小初 / 副标题黑体二号 / 学校·团队·指导老师·日期黑体小三，全部居中。"""
    # 把封面文字压到页面纵向约 1/3 处
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(150)
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing = 1.0
    _font(sp.add_run(''), 1)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.3
    _font(p.add_run(title or ''), SZ_XIAOCHU, CN_HEI, bold=True)

    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.3
        _font(p.add_run(subtitle), SZ_ER, CN_HEI)

    # 学校 / 团队 / 指导老师：黑体小三居中（申报书封面必备信息）
    info_lines = [x for x in (org, school, team, advisor) if x]
    for line in info_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.5
        _font(p.add_run(line), SZ_XIAOSAN, CN_HEI)

    if info_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(36)
        p.paragraph_format.line_spacing = 1.0
        _font(p.add_run(''), SZ_XIAOSAN)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    _font(p.add_run(datetime.datetime.now().strftime('%Y 年 %m 月 %d 日')),
          SZ_XIAOSAN, CN_HEI)


def _add_bookmark(p, name, bid):
    """给段落加书签，供目录的 PAGEREF 域取页码。"""
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bid))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bid))
    p._p.insert(0, start)
    p._p.append(end)


def _add_toc(doc, headings):
    """目录页：逐条列出真实标题 + PAGEREF 页码域。

    不用 TOC 域的原因：TOC 域必须用户手动「更新域」才出页码，
    很多人打开看到的是一片空白。这里改成直接写条目文字（一定可见），
    页码用 PAGEREF 域（dirty=true，Word 打开时自动填真实页码）。
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    _font(p.add_run('目  录'), SZ_ER, CN_HEI, bold=True)

    for level, text, mark in headings:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.0 if level == 1 else (0.6 if level == 2 else 1.2))
        pf.space_after = Pt(3)
        pf.line_spacing = 1.4

        # 右对齐制表位 + 省略号前导符，页码落在行尾
        pPr = p._p.get_or_add_pPr()
        tabs = OxmlElement('w:tabs')
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'right')
        tab.set(qn('w:leader'), 'dot')
        tab.set(qn('w:pos'), '9060')      # 可编辑区右边界 ≈16cm
        tabs.append(tab)
        pPr.append(tabs)

        _font(p.add_run(_fullwidth(text)), SZ_XIAOSI,
              CN_HEI if level == 1 else CN_SONG, bold=(level == 1))

        # 制表符
        tr = p.add_run()
        tr._r.append(OxmlElement('w:tab'))

        # PAGEREF 页码域
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'PAGEREF %s \\h' % mark)
        fld.set(qn('w:dirty'), 'true')
        rr = OxmlElement('w:r')
        tt = OxmlElement('w:t')
        tt.text = '1'
        rr.append(tt)
        fld.append(rr)
        p._p.append(fld)


# ============ 各类块渲染 ============

def _render_heading(doc, level, text, prof=None, mark=None, bid=0):
    """标题：按文档类型档位取字号/字体（申报书=标准层级，大纲=小字号轻层级）。"""
    prof = prof or PROFILES['default']
    size, cn, bold = prof['headings'].get(level, prof['headings'][4])
    # 必须挂 Heading 样式，否则目录域抓不到这一行
    try:
        p = doc.add_paragraph(style='Heading %d' % level)
    except KeyError:
        p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(14 if level == 1 else 10)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    pf.keep_with_next = True
    _font(p.add_run(_fullwidth(text)), size, cn, bold=bold)
    if mark:
        _add_bookmark(p, mark, bid)


def _render_para(doc, lines, prof=None):
    """正文：宋体小四，1.5 倍行距，段前段后 0。
    申报书/分析/报告 → 首行缩进 2 字符；大纲 → 不缩进、靠左对齐。
    行内的 **加粗** 保留（仅关键词，不通篇加粗）。"""
    prof = prof or PROFILES['default']
    _INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = prof['body_align']
        if prof['body_indent']:
            _indent_chars(p, 200)
        for seg in _INLINE.split(line):
            if not seg:
                continue
            if len(seg) > 4 and seg.startswith('**') and seg.endswith('**'):
                _font(p.add_run(_fullwidth(seg[2:-2])), SZ_XIAOSI, CN_SONG, bold=True)
            elif len(seg) > 2 and seg.startswith('`') and seg.endswith('`'):
                _font(p.add_run(seg[1:-1]), SZ_XIAOSI, EN)
            else:
                _font(p.add_run(_fullwidth(seg)), SZ_XIAOSI, CN_SONG)


def _render_list(doc, items, ordered, prof=None):
    """手动编号，不使用 Word 自动编号（规范要求）。"""
    prof = prof or PROFILES['default']
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = prof['body_align']
        _hanging_indent(p, 200)
        if ordered:
            _font(p.add_run('%d. ' % idx), SZ_XIAOSI, CN_SONG, bold=True)
        else:
            _font(p.add_run('● '), SZ_XIAOSI, CN_SONG)
        _render_inline_into(p, item)


def _render_inline_into(p, text):
    _INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')
    for seg in _INLINE.split(text):
        if not seg:
            continue
        if len(seg) > 4 and seg.startswith('**') and seg.endswith('**'):
            _font(p.add_run(_fullwidth(seg[2:-2])), SZ_XIAOSI, CN_SONG, bold=True)
        elif len(seg) > 2 and seg.startswith('`') and seg.endswith('`'):
            _font(p.add_run(seg[1:-1]), SZ_XIAOSI, EN)
        else:
            _font(p.add_run(_fullwidth(seg)), SZ_XIAOSI, CN_SONG)


def _render_quote(doc, lines):
    """引用/注释：宋体五号，左右缩进，无底色无边框。"""
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.8)
        pf.right_indent = Cm(0.4)
        pf.line_spacing = 1.5
        pf.space_before = Pt(3)
        pf.space_after = Pt(3)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _font(p.add_run(_fullwidth(line)), SZ_WU, CN_SONG)


def _render_code(doc, lines):
    """代码/公式块：Times New Roman 五号，缩进，无底色。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.8)
    pf.line_spacing = 1.25
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        _font(p.add_run(line if line else ' '), SZ_WU, EN)


def _render_caption(doc, text):
    """表题/图题：宋体五号加粗居中。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(8)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.5
    _font(p.add_run(_fullwidth(text)), SZ_WU, CN_SONG, bold=True)


def _render_table(doc, header, rows):
    """表格：单细线框，无底纹无特效；表头加粗居中，表内宋体小四。"""
    cols = len(header)
    if cols == 0:
        return
    norm = [(r + [''] * cols)[:cols] for r in rows]

    table = doc.add_table(rows=1, cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_cell_margins(table)

    hdr = table.rows[0]
    _repeat_header(hdr)
    for j, txt in enumerate(header):
        cell = hdr.cells[j]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(0)
        _font(p.add_run(_fullwidth(_plain(txt))), SZ_XIAOSI, CN_HEI, bold=True)

    for r in norm:
        row = table.add_row()
        for j, txt in enumerate(r):
            cell = row.cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            _font(p.add_run(_fullwidth(_plain(txt))), SZ_XIAOSI, CN_SONG)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)
    sp.paragraph_format.line_spacing = 1.0


def _render_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0


# ============ 主入口 ============

def build_docx(title, text, subtitle=None, org=None, toc=None,
               cover=None, cover_image=False, topic=None, doc_type='report',
               school=None, team=None, advisor=None):
    """生成符合格式标准的 Word 文档。

    :param title:       文档主标题（封面黑体小初）
    :param text:        Markdown 正文
    :param subtitle:    封面副标题（黑体二号），如赛事名称 / 项目类别
    :param org:         封面单位行（黑体小三），如「XX大学 · 团队名称」
    :param toc:         是否插入目录页；None = 按 doc_type 档位决定
    :param cover:       是否生成封面页；None = 按 doc_type 档位决定
    :param cover_image: 是否加封面背景图 —— 默认 False，
                        因为申报书标准禁止底色/特效，仅非申报场景可开
    :param topic:       封面背景图风格匹配用文本（cover_image=True 时生效）
    :param doc_type:    'report' 报告（默认）/ 'analysis' 分析 / 'outline' 大纲
                        / 'default' 申报书 —— 决定标题字号、正文缩进、封面目录
    :return:            python-docx Document
    """
    prof = PROFILES.get(doc_type, PROFILES['default'])
    if cover is None:
        cover = prof['cover']
    if toc is None:
        toc = prof['toc']

    doc = Document()
    _setup_page(doc, prof)
    _page_number_footer(doc)     # 页脚：页码居中（封面因首页不同自动无页码）
    _header(doc, title)          # 页眉：文档标题，宋体五号居中

    if cover:
        if cover_image:
            try:
                from cover_style import get_cover
                path, _sk, _ink = get_cover(topic or subtitle or title or '')
                _cover_background(doc, path)
            except Exception as e:
                print('[docx_render] 封面背景图失败，忽略：%r' % (e,))
        _cover(doc, title, subtitle, org, school, team, advisor)
        doc.add_page_break()

    blocks = _blocks(text or '')
    # 先收集标题，用来生成目录（目录条目必须和正文一一对应）
    headings, _bm_id = [], 0
    for blk in blocks:
        if blk[0] == 'h' and blk[1] <= 3:
            _bm_id += 1
            headings.append((blk[1], blk[2], '_toc_%d' % _bm_id))

    if toc:
        if headings:
            _add_toc(doc, headings)
        doc.add_page_break()

    _bm_i = 0
    for i, blk in enumerate(blocks):
        kind = blk[0]
        if kind == 'h':
            _bm_i += 1
            _render_heading(doc, blk[1], blk[2], prof,
                            mark=('_toc_%d' % _bm_i if blk[1] <= 3 else None),
                            bid=1000 + _bm_i)
        elif kind == 'p':
            # 形如「表 5-1 xxx」「图 3-2 xxx」且后面紧跟表格 → 按表题/图题渲染
            nxt = blocks[i + 1] if i + 1 < len(blocks) else None
            body = ' '.join(blk[1])
            if nxt and nxt[0] == 'table' and _RE_CAPTION.match(body) and len(body) <= 40:
                _render_caption(doc, body)
            elif _RE_CAPTION.match(body) and len(body) <= 40:
                _render_caption(doc, body)
            else:
                _render_para(doc, blk[1], prof)
        elif kind == 'ul':
            _render_list(doc, blk[1], ordered=False, prof=prof)
        elif kind == 'ol':
            _render_list(doc, blk[1], ordered=True, prof=prof)
        elif kind == 'quote':
            _render_quote(doc, blk[1])
        elif kind == 'code':
            _render_code(doc, blk[1])
        elif kind == 'table':
            _render_table(doc, blk[1], blk[2])
        elif kind == 'hr':
            _render_hr(doc)

    cp = doc.core_properties
    cp.title = title
    cp.author = '科创赛事多智能体协同创作助手'
    cp.category = '赛事申报材料'
    return doc


# ---- 封面背景图（默认不启用，保留给非申报场景）----

_NS_DRAW = (
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)


def _add_float_picture(paragraph, img_path, part, width_cm, height_cm,
                       pos_x_cm=0.0, pos_y_cm=0.0):
    from docx.oxml import parse_xml
    _img_part, r_id = part.get_or_add_image(img_path)
    cx = int(width_cm * 360000)
    cy = int(height_cm * 360000)
    xml = (
        '<w:drawing {ns}>'
        '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0"'
        ' relativeHeight="251658240" behindDoc="1" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>{px}</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>{py}</wp:posOffset></wp:positionV>'
        '<wp:extent cx="{cx}" cy="{cy}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:wrapNone/>'
        '<wp:docPr id="{pid}" name="cover"/>'
        '<a:graphic>'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic>'
        '<pic:nvPicPr><pic:cNvPr id="{pid}" name="cover"/><pic:cNvPicPr><a:picLocks/></pic:cNvPicPr></pic:nvPicPr>'
        '<pic:blipFill><a:blip r:embed="{rid}"/><a:srcRect/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        '<pic:spPr bwMode="auto">'
        '<a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:noFill/><a:ln><a:noFill/></a:ln>'
        '</pic:spPr>'
        '</pic:pic>'
        '</a:graphicData>'
        '</a:graphic>'
        '</wp:anchor>'
        '</w:drawing>'
    ).format(ns=_NS_DRAW, px=int(pos_x_cm * 360000), py=int(pos_y_cm * 360000),
             cx=cx, cy=cy, pid=abs(hash(img_path)) % 100000 + 1, rid=r_id)
    run = paragraph.add_run()
    run._r.append(parse_xml(xml))
    return run


def _cover_background(doc, img_path):
    header = doc.sections[0].first_page_header
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    _add_float_picture(p, img_path, header.part, 21.0, 29.7, 0, 0)
