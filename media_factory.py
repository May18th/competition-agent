# -*- coding: utf-8 -*-
"""
media_factory.py —— 媒体工厂（科创赛事助手）

把深度版流水线产出的纯文本，变成三样「真东西」：
  1. 真表格 —— Word / PPT 里的 table 对象，不是 markdown 竖线
  2. 真图表 —— matplotlib 出图，数据全部来自项目本身，不是写死的占位数据
  3. 真 PPT —— .pptx 路演稿，不是「导出成 docx 的假 PPT」

设计原则
  - 不依赖 Flask，纯函数，可以单独 import 测试
  - 图表统一白底输出：Word 里直接插，PPT 里放白卡片上，一套图两处用
  - 任何一步失败都不许炸掉主流程，返回降级结果由调用方决定

作者：WorkBuddy（阿渡）  2026-09-25
"""
import io
import os
import re
import math

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches as DocxInches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

from pptx import Presentation
from pptx.util import Inches, Pt as PPt, Emu
from pptx.dml.color import RGBColor as PRGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


# ==========================================================================
# 0. 字体与配色
# ==========================================================================
_FONT_DIR = "C:/Windows/Fonts"
for _f in ("msyh.ttc", "msyhbd.ttc", "simhei.ttf", "simsun.ttc"):
    _p = os.path.join(_FONT_DIR, _f)
    if os.path.exists(_p):
        try:
            font_manager.fontManager.addfont(_p)
        except Exception:
            pass

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "SimSun", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["figure.facecolor"] = "white"
plt.rcParams["axes.facecolor"] = "white"
plt.rcParams["savefig.facecolor"] = "white"

CN_FONT = "微软雅黑"

# 图表配色（白底）
INK = "#0f172a"
SUB = "#475569"
GRID = "#e2e8f0"
ACCENT = "#6366f1"
CYAN = "#22d3ee"
PURPLE = "#c084fc"
GREEN = "#34d399"
AMBER = "#fbbf24"
ROSE = "#fb7185"
PALETTE = [ACCENT, CYAN, PURPLE, GREEN, AMBER, ROSE, "#60a5fa", "#a3e635"]

# PPT 深色主题（与网页 WB-TECH 主题一致）
PPT_BG = PRGBColor(0x0a, 0x0f, 0x2b)
PPT_BG_ALT = PRGBColor(0x12, 0x1a, 0x3d)
PPT_TEXT = PRGBColor(0xe9, 0xee, 0xfc)
PPT_DIM = PRGBColor(0x93, 0xa0, 0xc4)
PPT_ACCENT = PRGBColor(0x81, 0x8c, 0xf8)
PPT_CARD = PRGBColor(0xff, 0xff, 0xff)


def _to_float(v, default=0.0):
    """把模型给的乱七八糟数字（"12万"/"12%"/"约30"）尽力转成 float"""
    if v is None:
        return default
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    s = re.sub(r"[，,\s]", "", s)
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        return default
    val = float(m.group(0))
    if "亿" in s:
        val *= 10000
    elif "万" in s:
        val *= 1
    return val


def _short(s, n):
    s = str(s or "").strip()
    return s if len(s) <= n else s[: n - 1] + "…"


# ==========================================================================
# 1. 图表
# ==========================================================================
def _finish(fig, width=9.0, height=4.8, dpi=170):
    fig.set_size_inches(width, height)
    try:
        fig.tight_layout()
    except Exception:
        pass
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight", pad_inches=0.18)
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def _style_axes(ax):
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.spines["left"].set_color(GRID)
    ax.spines["bottom"].set_color(GRID)
    ax.tick_params(colors=SUB, labelsize=10)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8, linestyle="--", alpha=0.9)
    ax.set_axisbelow(True)


def chart_score_radar(dims, values, title="项目综合评分雷达图"):
    """评分雷达图。dims: ['创新性','可行性',...]  values: 同长 0-100"""
    dims = [str(d) for d in (dims or [])]
    values = [_to_float(v, 0) for v in (values or [])]
    n = min(len(dims), len(values))
    if n < 3:
        return None
    dims, values = dims[:n], values[:n]
    values = [max(0.0, min(100.0, v)) for v in values]

    ang = [i * 2 * math.pi / n for i in range(n)]
    ang += ang[:1]
    vals = values + values[:1]

    fig = plt.figure()
    ax = fig.add_subplot(111, projection="polar")
    ax.plot(ang, vals, color=ACCENT, linewidth=2.2)
    ax.fill(ang, vals, color=ACCENT, alpha=0.22)
    ax.set_xticks(ang[:-1])
    ax.set_xticklabels(dims, fontsize=11, color=INK)
    ax.set_ylim(0, 100)
    ax.set_yticks([20, 40, 60, 80, 100])
    ax.set_yticklabels(["20", "40", "60", "80", "100"], fontsize=8, color=SUB)
    ax.grid(color=GRID, linewidth=0.8)
    ax.spines["polar"].set_color(GRID)
    ax.set_title(title, fontsize=14, color=INK, pad=18, fontweight="bold")

    for i, v in enumerate(values):
        ax.annotate(f"{v:.0f}", (ang[i], v), textcoords="offset points",
                    xytext=(6, 6), fontsize=9, color=ACCENT, fontweight="bold")
    return _finish(fig, 7.2, 5.4)


def chart_competitor(items, metrics, matrix, title="竞品能力对比"):
    """竞品对比分组柱状图。
    items:   ['本项目','竞品A','竞品B']（第一个高亮为「我们」）
    metrics: ['功能完整度','价格优势',...]
    matrix:  [[..每 metric 的得分..], ...]  与 items 同序，0-10 分
    """
    items = [str(i) for i in (items or [])]
    metrics = [str(m) for m in (metrics or [])]
    if not items or not metrics or len(matrix or []) < len(items):
        return None
    matrix = [row[: len(metrics)] for row in matrix[: len(items)]]
    matrix = [[_to_float(x, 0) for x in row] for row in matrix]

    n_m = len(metrics)
    n_i = len(items)
    x = list(range(n_m))
    width = min(0.8 / max(n_i, 1), 0.22)

    fig, ax = plt.subplots()
    for i in range(n_i):
        offset = (i - (n_i - 1) / 2) * width
        vals = [matrix[i][j] for j in range(n_m)]
        color = PALETTE[0] if i == 0 else PALETTE[(i % (len(PALETTE) - 1)) + 1]
        alpha = 1.0 if i == 0 else 0.72
        bars = ax.bar([p + offset for p in x], vals, width, label=items[i],
                      color=color, alpha=alpha)
        if i == 0:
            for b, v in zip(bars, vals):
                ax.annotate(f"{v:.1f}", (b.get_x() + b.get_width() / 2, v),
                            textcoords="offset points", xytext=(0, 4),
                            ha="center", fontsize=8.5, color=ACCENT, fontweight="bold")

    ax.set_xticks(x)
    ax.set_xticklabels([_short(m, 10) for m in metrics], fontsize=11, color=INK)
    ax.set_ylabel("评分（满分10）", fontsize=10, color=SUB)
    ax.set_title(title, fontsize=14, color=INK, fontweight="bold", pad=12)
    ax.legend(fontsize=9, frameon=False, ncol=min(n_i, 4))
    _style_axes(ax)
    ax.set_ylim(0, 11)
    return _finish(fig, 9.0, 4.8)


def chart_finance(years, revenue, cost, profit, unit="万元", title="三年财务预测"):
    """三年财务预测：收入/成本柱 + 利润折线"""
    years = [str(y) for y in (years or [])]
    revenue = [_to_float(v) for v in (revenue or [])]
    cost = [_to_float(v) for v in (cost or [])]
    profit = [_to_float(v) for v in (profit or [])]
    n = min(len(years), len(revenue), len(cost))
    if n == 0:
        return None
    years, revenue, cost = years[:n], revenue[:n], cost[:n]
    profit = profit[:n] if profit else [revenue[i] - cost[i] for i in range(n)]

    x = list(range(n))
    w = 0.34
    fig, ax = plt.subplots()
    ax.bar([p - w / 2 for p in x], revenue, w, label="营业收入", color=ACCENT)
    ax.bar([p + w / 2 for p in x], cost, w, label="总成本", color=PURPLE, alpha=0.85)
    ax.plot(x, profit, marker="o", linewidth=2.2, color=GREEN, label="净利润")

    for i, v in enumerate(profit):
        ax.annotate(f"{v:.0f}", (i, v), textcoords="offset points",
                    xytext=(0, 8), ha="center", fontsize=9,
                    color=GREEN, fontweight="bold")

    ax.set_xticks(x)
    ax.set_xticklabels(years, fontsize=11, color=INK)
    ax.set_ylabel(f"金额（{unit}）", fontsize=10, color=SUB)
    ax.set_title(title, fontsize=14, color=INK, fontweight="bold", pad=12)
    ax.legend(fontsize=9, frameon=False)
    _style_axes(ax)
    return _finish(fig, 9.0, 4.8)


def chart_timeline(phases, start_week, end_week, title="项目实施计划（甘特图）"):
    """实施计划甘特图。phases/start_week/end_week 同序，周为单位"""
    phases = [str(p) for p in (phases or [])]
    starts = [_to_float(s) for s in (start_week or [])]
    ends = [_to_float(e) for e in (end_week or [])]
    n = min(len(phases), len(starts), len(ends))
    if n == 0:
        return None
    phases, starts, ends = phases[:n], starts[:n], ends[:n]
    for i in range(n):
        if ends[i] <= starts[i]:
            ends[i] = starts[i] + 1

    fig, ax = plt.subplots()
    ypos = list(range(n))[::-1]
    colors = [PALETTE[i % len(PALETTE)] for i in range(n)]
    ax.barh(ypos, [ends[i] - starts[i] for i in range(n)],
            left=starts, height=0.55, color=colors, alpha=0.92)
    for i, y in enumerate(ypos):
        ax.text(starts[i] + (ends[i] - starts[i]) / 2, y,
                f"W{int(starts[i]) + 1}-W{int(ends[i])}",
                va="center", ha="center", fontsize=9, color="white", fontweight="bold")

    ax.set_yticks(ypos)
    ax.set_yticklabels([_short(p, 14) for p in phases], fontsize=11, color=INK)
    ax.set_xlabel("周次", fontsize=10, color=SUB)
    ax.set_title(title, fontsize=14, color=INK, fontweight="bold", pad=12)
    total = max(max(ends), 1)
    ax.set_xlim(0, total + 0.6)
    ax.set_xticks(list(range(0, int(total) + 2)))
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.spines["left"].set_color(GRID)
    ax.spines["bottom"].set_color(GRID)
    ax.tick_params(colors=SUB, labelsize=10)
    ax.xaxis.grid(True, color=GRID, linewidth=0.8, linestyle="--", alpha=0.9)
    ax.set_axisbelow(True)
    return _finish(fig, 9.0, max(3.0, 1.0 + n * 0.62))


def chart_budget(labels, values, title="预算构成"):
    """预算构成环形图"""
    labels = [str(l) for l in (labels or [])]
    values = [_to_float(v, 0) for v in (values or [])]
    n = min(len(labels), len(values))
    if n == 0 or sum(values[:n]) <= 0:
        return None
    labels, values = labels[:n], values[:n]

    fig, ax = plt.subplots(figsize=(7.4, 5.0))
    wedges, texts, autotexts = ax.pie(
        values, labels=[_short(l, 12) for l in labels], autopct=lambda p: f"{p:.0f}%",
        startangle=100, colors=PALETTE[:n],
        wedgeprops={"width": 0.42, "edgecolor": "white", "linewidth": 2},
        textprops={"fontsize": 11, "color": INK},
        pctdistance=0.78)
    for t in autotexts:
        t.set_fontsize(9.5)
        t.set_color("white")
        t.set_fontweight("bold")
    ax.set_title(title, fontsize=14, color=INK, fontweight="bold", pad=10)
    ax.axis("equal")
    return _finish(fig, 7.4, 5.0)


def chart_barh(labels, values, title="对比", xlabel="", color=ACCENT):
    """通用横向柱状图（兜底图）"""
    labels = [str(l) for l in (labels or [])]
    values = [_to_float(v, 0) for v in (values or [])]
    n = min(len(labels), len(values))
    if n == 0:
        return None
    labels, values = labels[:n], values[:n]
    fig, ax = plt.subplots()
    ypos = list(range(n))[::-1]
    ax.barh(ypos, values, height=0.55, color=color, alpha=0.9)
    for y, v in zip(ypos, values):
        ax.text(v, y, f" {v:g}", va="center", fontsize=9.5, color=INK, fontweight="bold")
    ax.set_yticks(ypos)
    ax.set_yticklabels([_short(l, 14) for l in labels], fontsize=11, color=INK)
    ax.set_xlabel(xlabel, fontsize=10, color=SUB)
    ax.set_title(title, fontsize=14, color=INK, fontweight="bold", pad=12)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    ax.spines["left"].set_color(GRID)
    ax.spines["bottom"].set_color(GRID)
    ax.tick_params(colors=SUB, labelsize=10)
    ax.xaxis.grid(True, color=GRID, linewidth=0.8, linestyle="--", alpha=0.9)
    ax.set_axisbelow(True)
    return _finish(fig, 9.0, max(3.0, 1.0 + n * 0.62))


# ==========================================================================
# 2. Markdown 表格解析 —— 把竖线变成真表格
# ==========================================================================
_ROW_RE = re.compile(r"^\s*\|(.+)\|\s*$")
_SEP_RE = re.compile(r"^\s*\|?[\s:\-\|]+\|[\s:\-\|]*$")


def _split_row(line):
    inner = _ROW_RE.match(line).group(1)
    return [c.strip() for c in inner.split("|")]


def split_blocks(text):
    """把整段文本切成 [('p', 文本块), ('table', 二维表)] 序列"""
    blocks = []
    if not text:
        return blocks
    buf = []
    table = []

    def flush_buf():
        if buf:
            blocks.append(("p", "\n".join(buf).strip()))
            buf.clear()

    def flush_table():
        if table:
            blocks.append(("table", [r[:] for r in table]))
            table.clear()

    for raw in str(text).split("\n"):
        line = raw.rstrip()
        if _ROW_RE.match(line):
            if _SEP_RE.match(line) and table:
                continue
            flush_buf()
            table.append(_split_row(line))
        else:
            flush_table()
            buf.append(line)
    flush_table()
    flush_buf()
    # 只有表头没有数据的表不要
    return [b for b in blocks if not (b[0] == "table" and len(b[1]) < 2)]


def count_tables(text):
    return sum(1 for b in split_blocks(text or "") if b[0] == "table")


# ==========================================================================
# 3. Word：带真表格、可插图的 docx
# ==========================================================================
def _clean_inline(s):
    s = re.sub(r"\*\*(.*?)\*\*", r"\1", s)
    s = re.sub(r"(?<!\*)\*(?!\*)(.*?)(?<!\*)\*(?!\*)", r"\1", s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    return s.strip()


def _set_run(run, size=11, bold=False, color=None, name=CN_FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    # 中文字体必须再设一次东亚字体，否则 Word 里显示成宋体
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _add_docx_table(doc, rows):
    """往 docx 里塞一张真表格"""
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(_clean_inline(str(val)))
            _set_run(run, size=10, bold=(ri == 0),
                     color=RGBColor(0xff, 0xff, 0xff) if ri == 0 else RGBColor(0x1f, 0x29, 0x37))
            if ri == 0:
                shd = cell._tc.get_or_add_tcPr()
                from docx.oxml import OxmlElement
                el = OxmlElement("w:shd")
                el.set(qn("w:val"), "clear")
                el.set(qn("w:fill"), "6366F1")
                shd.append(el)
            elif ri % 2 == 0:
                from docx.oxml import OxmlElement
                shd = cell._tc.get_or_add_tcPr()
                el = OxmlElement("w:shd")
                el.set(qn("w:val"), "clear")
                el.set(qn("w:fill"), "F1F5FF")
                shd.append(el)
    doc.add_paragraph()


def build_rich_docx(title, text, chart_items=None, subtitle=None):
    """生成带真表格的 docx。
    chart_items: [(png_bytes, 图注)]，会插在正文末尾
    返回 bytes
    """
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = CN_FONT
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    for section in doc.sections:
        section.top_margin = Cm(2.4)
        section.bottom_margin = Cm(2.4)
        section.left_margin = Cm(2.6)
        section.right_margin = Cm(2.6)

    h = doc.add_heading("", 0)
    run = h.add_run(title)
    _set_run(run, size=20, bold=True, color=RGBColor(0x4f, 0x46, 0xe5))
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle:
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = sp.add_run(subtitle)
        _set_run(r, size=10, color=RGBColor(0x64, 0x74, 0x8b))
    doc.add_paragraph()

    for kind, payload in split_blocks(text or ""):
        if kind == "table":
            _add_docx_table(doc, payload)
            continue
        for line in payload.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("#"):
                level = min(line.count("#"), 4)
                p = doc.add_heading("", level=level)
                r = p.add_run(_clean_inline(line.lstrip("#").strip()))
                _set_run(r, size={1: 16, 2: 14}.get(level, 12), bold=True,
                         color=RGBColor(0x4a, 0x55, 0xb8))
            elif line.startswith(("- ", "• ", "· ", "* ")):
                p = doc.add_paragraph(style="List Bullet")
                r = p.add_run(_clean_inline(line[2:]))
                _set_run(r, size=11)
                p.paragraph_format.line_spacing = 1.5
            else:
                p = doc.add_paragraph()
                r = p.add_run(_clean_inline(line))
                _set_run(r, size=11)
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.first_line_indent = Cm(0.74)
                p.paragraph_format.space_after = Pt(6)

    for png, caption in (chart_items or []):
        if not png:
            continue
        doc.add_paragraph()
        try:
            doc.add_picture(io.BytesIO(png), width=DocxInches(5.6))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            continue
        if caption:
            cp = doc.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cp.add_run(caption)
            _set_run(r, size=9, color=RGBColor(0x64, 0x74, 0x8b))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ==========================================================================
# 4. PPT：真 .pptx 路演稿
# ==========================================================================
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
M = Inches(0.78)          # 左右边距
CONTENT_W = Inches(11.77)


def _ppt_run(tf, text, size=16, bold=False, color=PPT_TEXT, space_after=8,
             align=PP_ALIGN.LEFT, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = PPt(space_after)
    r = p.add_run()
    r.text = text
    r.font.size = PPt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = CN_FONT
    return p


def _new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = PPT_BG
    return slide


def _deco_bar(slide, color=PPT_ACCENT, top=Inches(0.62), h=Inches(0.42), w=Inches(0.09)):
    shp = slide.shapes.add_shape(1, M, top, w, h)   # 1 = MSO_SHAPE.RECTANGLE
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _footer(slide, idx, total, brand):
    tf = slide.shapes.add_textbox(M, Inches(6.92), Inches(9.0), Inches(0.32)).text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].runs
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = _short(brand, 46)
    run.font.size = PPt(9)
    run.font.color.rgb = PPT_DIM
    run.font.name = CN_FONT

    tf2 = slide.shapes.add_textbox(Inches(11.6), Inches(6.92), Inches(0.95), Inches(0.32)).text_frame
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = f"{idx} / {total}"
    r2.font.size = PPt(9)
    r2.font.color.rgb = PPT_DIM
    r2.font.name = CN_FONT


def _add_title(slide, title, kicker=None):
    if kicker:
        tb = slide.shapes.add_textbox(M + Inches(0.22), Inches(0.58), CONTENT_W, Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        _ppt_run(tf, str(kicker), size=11, bold=True, color=PPT_ACCENT, first=True, space_after=0)
    _deco_bar(slide, top=Inches(0.92))
    tb = slide.shapes.add_textbox(M + Inches(0.22), Inches(0.95), CONTENT_W, Inches(0.72))
    tf = tb.text_frame
    tf.word_wrap = True
    _ppt_run(tf, _short(title, 40), size=27, bold=True, color=PPT_TEXT, first=True, space_after=0)


def _slide_bullets(prs, slide, bullets, idx, total, brand):
    n = len(bullets)
    size = 17 if n <= 5 else (15 if n <= 7 else 13)
    top = Inches(1.95)
    tb = slide.shapes.add_textbox(M + Inches(0.22), top, CONTENT_W, Inches(4.85))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for b in bullets:
        b = _clean_inline(str(b)).strip()
        if not b:
            continue
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = PPt(14 if n <= 6 else 9)
        r = p.add_run()
        r.text = "▍ " + b
        r.font.size = PPt(size)
        r.font.color.rgb = PPT_TEXT
        r.font.name = CN_FONT
    _footer(slide, idx, total, brand)


def _slide_chart(prs, slide, png, caption=None, idx=1, total=1, brand=""):
    """图表页：白卡片 + 图"""
    card = slide.shapes.add_shape(1, Inches(0.92), Inches(1.85), Inches(11.5), Inches(4.75))
    card.fill.solid()
    card.fill.fore_color.rgb = PPT_CARD
    card.line.fill.background()
    card.shadow.inherit = False

    try:
        from PIL import Image as _Img
        im = _Img.open(io.BytesIO(png))
        ratio = im.height / im.width
    except Exception:
        ratio = 0.55

    max_w = Inches(10.7)
    max_h = Inches(4.15)
    w = max_w
    h = int(w * ratio)
    if h > max_h:
        h = max_h
        w = int(h / ratio) if ratio else max_w
    left = Inches(0.92) + int((Inches(11.5) - w) / 2)
    top = Inches(1.85) + int((Inches(4.75) - h) / 2)
    slide.shapes.add_picture(io.BytesIO(png), left, top, w, h)

    if caption:
        tb = slide.shapes.add_textbox(Inches(0.92), Inches(6.72), Inches(11.5), Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = _short(caption, 60)
        r.font.size = PPt(10)
        r.font.color.rgb = PPT_DIM
        r.font.name = CN_FONT
    _footer(slide, idx, total, brand)


def _slide_table(prs, slide, rows, idx=1, total=1, brand=""):
    rows = [[_clean_inline(str(c)) for c in r] for r in (rows or []) if any(str(c).strip() for c in r)]
    if not rows:
        return
    nrow = min(len(rows), 9)
    ncol = max(len(r) for r in rows[:nrow])
    rows = [r + [""] * (ncol - len(r)) for r in rows[:nrow]]

    top = Inches(1.92)
    height = Inches(0.42) * nrow + Inches(0.1)
    shp = slide.shapes.add_table(nrow, ncol, Inches(0.92), top, Inches(11.5), height)
    tbl = shp.table

    # 列宽：按内容长度加权
    weights = []
    for ci in range(ncol):
        weights.append(max(len(rows[ri][ci]) for ri in range(nrow)) + 2)
    total_w = sum(weights)
    for ci in range(ncol):
        tbl.columns[ci].width = Emu(int(Inches(11.5) * weights[ci] / total_w))

    for ri in range(nrow):
        tbl.rows[ri].height = Inches(0.42)
        for ci in range(ncol):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = PRGBColor(0x63, 0x66, 0xf1)
            else:
                cell.fill.fore_color.rgb = PRGBColor(0x14, 0x1c, 0x42) if ri % 2 else PRGBColor(0x10, 0x16, 0x36)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if (ci == 0 or ri == 0) else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = _short(rows[ri][ci], 26)
            r.font.size = PPt(12 if ncol <= 4 else 10)
            r.font.bold = (ri == 0)
            r.font.name = CN_FONT
            r.font.color.rgb = PRGBColor(0xff, 0xff, 0xff) if ri == 0 else PPT_TEXT
    _footer(slide, idx, total, brand)


def build_pptx(deck, charts=None, out_path=None):
    """生成路演 PPT。
    deck = {
      "title": "项目名", "subtitle": "一句话定位", "brand": "页脚",
      "slides": [
        {"layout": "cover|section|bullets|chart|table|end",
         "title": "", "kicker": "01 / 痛点",
         "bullets": [], "note": "", "chart": "finance", "caption": "", "table": [[...]]}
      ]
    }
    charts: {key: png_bytes}
    返回 bytes
    """
    charts = charts or {}
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    title = deck.get("title") or "科创项目路演"
    subtitle = deck.get("subtitle") or ""
    brand = deck.get("brand") or title
    slides = deck.get("slides") or []

    # ---------- 封面 ----------
    cover = None
    for s in slides:
        if s.get("layout") == "cover":
            cover = s
            break
    cs = _new_slide(prs)
    band = cs.shapes.add_shape(1, 0, Inches(5.6), SLIDE_W, Inches(1.9))
    band.fill.solid()
    band.fill.fore_color.rgb = PRGBColor(0x12, 0x1a, 0x3d)
    band.line.fill.background()
    band.shadow.inherit = False

    tb = cs.shapes.add_textbox(Inches(1.0), Inches(2.15), Inches(11.3), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True
    _ppt_run(tf, _short(cover.get("title") or title, 34), size=42, bold=True,
             color=PPT_TEXT, first=True, space_after=6)
    if subtitle:
        _ppt_run(tf, _short(subtitle, 60), size=17, color=PRGBColor(0x67, 0xe8, 0xf9), space_after=0)

    tb2 = cs.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    _ppt_run(tf2, _short(brand, 70), size=12, color=PPT_DIM, first=True, space_after=0)

    # ---------- 正文页 ----------
    body = [s for s in slides if s.get("layout") != "cover"]
    total = len(body) + 1
    idx = 2
    for s in body:
        layout = s.get("layout") or "bullets"
        slide = _new_slide(prs)
        if layout == "end":
            tb = slide.shapes.add_textbox(Inches(1.2), Inches(3.0), Inches(10.9), Inches(1.6))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = _short(s.get("title") or "感谢聆听", 24)
            r.font.size = PPt(38)
            r.font.bold = True
            r.font.color.rgb = PPT_TEXT
            r.font.name = CN_FONT
            if s.get("bullets"):
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.CENTER
                r2 = p2.add_run()
                r2.text = _short(_clean_inline(str(s["bullets"][0])), 40)
                r2.font.size = PPt(15)
                r2.font.color.rgb = PRGBColor(0x67, 0xe8, 0xf9)
                r2.font.name = CN_FONT
            _footer(slide, idx, total, brand)
            idx += 1
            continue

        _add_title(slide, s.get("title") or "", kicker=s.get("kicker"))

        if layout == "chart":
            png = charts.get(s.get("chart"))
            if png:
                _slide_chart(prs, slide, png, s.get("caption"), idx, total, brand)
            else:
                _slide_bullets(prs, slide, s.get("bullets") or [], idx, total, brand)
        elif layout == "table":
            rows = s.get("table")
            if rows:
                _slide_table(prs, slide, rows, idx, total, brand)
            else:
                _slide_bullets(prs, slide, s.get("bullets") or [], idx, total, brand)
        else:
            _slide_bullets(prs, slide, s.get("bullets") or [], idx, total, brand)

        note = s.get("note")
        if note:
            try:
                slide.notes_slide.notes_text_frame.text = str(note)[:1800]
            except Exception:
                pass
        idx += 1

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    data = buf.read()
    if out_path:
        try:
            os.makedirs(os.path.dirname(out_path), exist_ok=True)
            with open(out_path, "wb") as f:
                f.write(data)
        except Exception:
            pass
    return data


# ==========================================================================
# 5. 自检
# ==========================================================================
if __name__ == "__main__":
    pngs = {}
    pngs["radar"] = chart_score_radar(["创新性", "可行性", "市场价值", "技术壁垒", "社会价值"],
                                      [82, 74, 68, 71, 80])
    pngs["competitor"] = chart_competitor(
        ["本项目", "竞品A", "竞品B"],
        ["功能完整度", "价格优势", "技术壁垒", "易用性"],
        [[9, 8, 9, 8], [7, 6, 5, 7], [6, 7, 4, 6]])
    pngs["finance"] = chart_finance(["第1年", "第2年", "第3年"], [120, 380, 860], [150, 300, 560], [-30, 80, 300])
    pngs["timeline"] = chart_timeline(["需求调研", "原型开发", "小批量试产", "市场推广"], [0, 2, 5, 7], [2, 5, 7, 8])
    pngs["budget"] = chart_budget(["研发", "生产", "营销", "运营"], [45, 30, 15, 10])
    for k, v in pngs.items():
        print(k, "OK" if v else "FAIL", len(v) if v else 0)

    deck = {
        "title": "智能危化品试剂管理柜",
        "subtitle": "让每一瓶试剂的取用都可追溯",
        "brand": "iCAN 大学生创新创业大赛 · 智能危化品试剂管理柜",
        "slides": [
            {"layout": "cover", "title": "智能危化品试剂管理柜"},
            {"layout": "bullets", "kicker": "01 / 痛点", "title": "实验室试剂管理有多痛",
             "bullets": ["500mL 丙酮领用后实际只用 80mL，余量无人追踪",
                         "纸质台账滞后 2-3 天，事故追溯靠翻本子",
                         "危化品超期存放普遍，盘点一次要 3 小时"],
             "note": "用具体场景开场，先讲痛点有多贵。"},
            {"layout": "chart", "kicker": "02 / 竞争力", "title": "与现有方案的差距",
             "chart": "competitor", "caption": "竞品能力对比（满分 10 分，数据来自本项目分析）"},
            {"layout": "table", "kicker": "03 / 竞品", "title": "竞品逐项对比",
             "table": [["方案", "单柜成本", "识别精度", "台账自动化"],
                       ["本项目", "1.2 万元", "±1 g", "全自动"],
                       ["传统人工台账", "0.1 万元", "人工", "无"],
                       ["进口智能柜", "6.8 万元", "±0.5 g", "半自动"]]},
            {"layout": "chart", "kicker": "04 / 财务", "title": "三年财务预测",
             "chart": "finance", "caption": "收入 / 成本 / 净利润（万元）"},
            {"layout": "chart", "kicker": "05 / 进度", "title": "八周落地计划",
             "chart": "timeline", "caption": "项目实施甘特图（周次）"},
            {"layout": "end", "title": "感谢聆听", "bullets": ["敬请评委老师指正"]},
        ],
    }
    out = build_pptx(deck, pngs, out_path=os.path.join(os.path.dirname(__file__), "_media_demo.pptx"))
    print("pptx bytes:", len(out))

    docx = build_rich_docx("测试文档", "## 竞品对比\n\n| 方案 | 成本 | 精度 |\n|---|---|---|\n| 本项目 | 1.2万 | ±1g |\n| 进口柜 | 6.8万 | ±0.5g |\n\n正文一段话。\n",
                           chart_items=[(pngs["finance"], "图：三年财务预测")])
    print("docx bytes:", len(docx), "表格数:", count_tables("| a | b |\n|---|---|\n| 1 | 2 |"))
