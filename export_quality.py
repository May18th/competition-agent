"""Hard-failure classification for generated deliverables."""
import re
import zipfile


def hard_content_issues(text, *, double_blind=False):
    value = str(text or "")
    issues = []
    if not value.strip():
        issues.append("正文为空，无法生成可交付文件")
        return issues
    if "\ufffd" in value or value.count("?") > max(3, len(value) * 0.3) or "？？" in value:
        issues.append("正文存在乱码，已暂停导出以避免交付损坏内容")
    if double_blind:
        if re.search(r"[\u4e00-\u9fff]{2,4}大学(?!生|城|名称|地址|简称|全称|校名)", value):
            issues.append("双盲评审材料仍含疑似院校名称")
        if re.search(
            r"(?:指导老师|指导教师|导师|负责人|队长)\s*[:：]\s*[\u4e00-\u9fff]{2,4}", value
        ):
            issues.append("双盲评审材料仍含疑似人员姓名")
    return issues


def hard_artifact_issues(kind, issues):
    """Select only corruption/empty/garbling from renderer diagnostics."""
    hard_markers = {
        "docx": ("docx 正文为空", "中文乱码", "docx 检查失败"),
        "pdf": ("PDF 无页面", "PDF 中文乱码", "pdf 检查失败"),
        "pptx": ("pptx 检查失败",),
        "zip": ("zip 检查失败",),
    }.get(kind, ())
    return [issue for issue in issues if any(marker in str(issue) for marker in hard_markers)]


def public_artifact_reasons(kind, issues):
    """Translate internal parser diagnostics into concise, non-technical notices."""
    if not issues:
        return []
    labels = {
        "docx": "Word 文件无法完整读取，系统已尝试自动修复",
        "pdf": "PDF 文件无法完整读取，系统已尝试自动修复",
        "pptx": "路演 PPT 文件无法完整读取，系统已尝试自动修复",
        "zip": "材料包结构异常，系统已尝试自动修复",
    }
    label = labels.get(kind, "导出文件无法完整读取，系统已尝试自动修复")
    if any("乱码" in str(issue) for issue in issues):
        return ["文件内容检测到乱码，系统已尝试自动修复"]
    if any("为空" in str(issue) or "无页面" in str(issue) for issue in issues):
        return ["文件内容为空，无法安全交付"]
    return [label]


def verify_artifact_integrity(kind, path):
    """Independent structural check used only if the full quality checker errors."""
    try:
        if kind in {"docx", "pptx", "zip"}:
            with zipfile.ZipFile(path) as archive:
                if archive.testzip() is not None:
                    return False
                names = set(archive.namelist())
                if kind == "docx":
                    required = {"[Content_Types].xml", "word/document.xml"}
                elif kind == "pptx":
                    required = {"[Content_Types].xml", "ppt/presentation.xml"}
                else:
                    return bool(names)
                if not required.issubset(names):
                    return False
                from xml.etree import ElementTree
                for member in required:
                    ElementTree.fromstring(archive.read(member))
            if kind == "docx":
                from docx import Document
                document = Document(path)
                return bool("".join(p.text for p in document.paragraphs).strip())
            if kind == "pptx":
                from pptx import Presentation
                return len(Presentation(path).slides) > 0
        elif kind == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(path, strict=True)
            return len(reader.pages) > 0
    except Exception:
        return False
    return False
